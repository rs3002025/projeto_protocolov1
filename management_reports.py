"""Indicadores gerenciais baseados em datas registradas, não em suposições."""
import io
from collections import defaultdict
from datetime import datetime,time,timedelta,timezone
from zoneinfo import ZoneInfo
from flask import abort,render_template,request,send_file
from sqlalchemy import and_,func,or_,select,type_coerce,DateTime

BR=ZoneInfo('America/Sao_Paulo')
CLOSED=('FINALIZADO','CONCLUÍDO','ARQUIVADO')


def utc(value):
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def management_data(args):
    from app import (accessible_protocols_query,current_tenant_id,db,parse_iso_date,tenant_query)
    from models import Protocolo,Movimentacao,ProvidenciaProtocolo,Lotacao
    from dashboard_workbench import today_brasilia
    end=parse_iso_date(args.get('data_fim'),'Data final') or today_brasilia()
    start=parse_iso_date(args.get('data_inicio'),'Data inicial') or end-timedelta(days=29)
    if start>end or (end-start).days>3660:
        abort(400,description='Informe um período válido de até dez anos.')
    begin=datetime.combine(start,time.min,tzinfo=BR).astimezone(timezone.utc).replace(tzinfo=None)
    finish=datetime.combine(end+timedelta(days=1),time.min,tzinfo=BR).astimezone(timezone.utc).replace(tzinfo=None)
    tenant=current_tenant_id(); query=accessible_protocols_query()
    entries=query.filter(Protocolo.data_solicitacao>=start,Protocolo.data_solicitacao<=end).with_entities(
        Protocolo.data_solicitacao,func.count(Protocolo.id)).group_by(Protocolo.data_solicitacao).all()
    # Datas UTC do banco são apresentadas no dia civil de Brasília.
    closing=query.filter(Protocolo.finalizado_em>=begin,Protocolo.finalizado_em<finish).with_entities(
        Protocolo.data_solicitacao,Protocolo.finalizado_em).all()
    daily=defaultdict(lambda:{'entradas':0,'conclusoes':0})
    for day,count in entries: daily[day]['entradas']=count
    valid_days=[]
    for requested,finished in closing:
        day=utc(finished).astimezone(BR).date(); daily[day]['conclusoes']+=1
        if requested and requested<=day: valid_days.append((day-requested).days)
    untracked=query.filter(Protocolo.finalizado_em.is_(None),or_(
        Protocolo.status.in_(['FINALIZADO','CONCLUÍDO']),Protocolo.arquivado_em.isnot(None),Protocolo.status=='ARQUIVADO')).count()
    accessible=query.with_entities(Protocolo.id).statement
    # A janela considera todo o histórico antes de recortar o período. Assim,
    # uma transferência posterior ao período encerra corretamente o atendimento.
    timeline=select(Movimentacao.id.label('id'),Movimentacao.providencia_id.label('participation_id'),
        func.max(Movimentacao.id).over(partition_by=Movimentacao.protocolo_id).label('latest_protocol_move'),
        Movimentacao.setor_destino_id.label('setor'),
        Movimentacao.enviado_em.label('sent'),Movimentacao.recebido_em.label('received'),
        type_coerce(func.lead(Movimentacao.enviado_em).over(partition_by=[Movimentacao.protocolo_id,Movimentacao.providencia_id],
            order_by=Movimentacao.id),DateTime()).label('next_sent'),
        ProvidenciaProtocolo.concluida_em.label('participation_end'),
        Protocolo.finalizado_em.label('closed_at'),Protocolo.arquivado_em.label('archived_at'),
        Protocolo.status.label('status')).join(Protocolo,Protocolo.id==Movimentacao.protocolo_id).outerjoin(
        ProvidenciaProtocolo,and_(ProvidenciaProtocolo.id==Movimentacao.providencia_id,
                                 ProvidenciaProtocolo.tenant_id==tenant)).where(
            Movimentacao.tenant_id==tenant,Protocolo.tenant_id==tenant,
            Movimentacao.protocolo_id.in_(accessible)).subquery('tempos_setoriais')
    movements=db.session.execute(select(timeline).where(timeline.c.sent>=begin,timeline.c.sent<finish)).mappings()
    names={s.id:s.nome for s in tenant_query(Lotacao).all()}
    sectors={}
    for movement in movements:
        row=sectors.setdefault(movement['setor'],dict(nome=names.get(movement['setor'],'Setor não informado'),
            enviados=0,recebidos=0,pendentes=0,redirecionados=0,sem_marco=0,wait=[],service=[]))
        row['enviados']+=1
        sent=utc(movement['sent']); received=utc(movement['received']) if movement['received'] else None
        endpoints=[utc(movement[k]) for k in ('next_sent','participation_end','closed_at','archived_at') if movement[k]]
        current=(not movement['next_sent'] and not movement['participation_end'] and
                 (movement['participation_id'] is not None or movement['id']==movement['latest_protocol_move']))
        if received:
            row['recebidos']+=1
            if received>=sent: row['wait'].append((received-sent).total_seconds()/3600)
            ends=[value for value in endpoints if value>=received]
            if ends: row['service'].append((min(ends)-received).total_seconds()/3600)
            elif movement['status'] in CLOSED or not current: row['sem_marco']+=1
        elif any(value>=sent for value in endpoints) or movement['status'] in CLOSED or not current:
            row['redirecionados']+=1
        else: row['pendentes']+=1
    for row in sectors.values():
        for values,target in (('wait','espera_media'),('service','atendimento_medio')):
            row[target]=round(sum(row[values])/len(row[values]),1) if row[values] else None
        row['base_espera']=len(row.pop('wait')); row['base_atendimento']=len(row.pop('service'))
    monthly=(end-start).days>92
    chart_totals=defaultdict(lambda:{'entradas':0,'conclusoes':0})
    for day,values in daily.items():
        key=day.replace(day=1) if monthly else day
        for kind in ('entradas','conclusoes'): chart_totals[key][kind]+=values[kind]
    chart=dict(labels=[],entradas=[],conclusoes=[],agrupamento='Por mês' if monthly else 'Por dia')
    cursor=start.replace(day=1) if monthly else start
    while cursor<=end:
        chart['labels'].append(cursor.strftime('%m/%Y' if monthly else '%d/%m/%Y'))
        for kind in ('entradas','conclusoes'): chart[kind].append(chart_totals[cursor][kind])
        cursor=(cursor.replace(day=28)+timedelta(days=4)).replace(day=1) if monthly else cursor+timedelta(days=1)
    return dict(start=start,end=end,entradas=sum(c for _,c in entries),conclusoes=len(closing),chart=chart,
        media_dias=round(sum(valid_days)/len(valid_days),1) if valid_days else None,base_media=len(valid_days),
        sem_data_conclusao=untracked,daily=[dict(data=d,**daily[d]) for d in sorted(daily)],
        sectors=sorted(sectors.values(),key=lambda r:(-r['pendentes'],-r['enviados'],r['nome'])))


def register_management_routes(app):
    from app import permission_required

    @app.get('/relatorios/gestao')
    @permission_required('reports')
    def relatorio_gestao():
        return render_template('relatorio_gestao.html',title='Indicadores de gestão',report=management_data(request.args))

    @app.get('/relatorios/gestao/excel')
    @permission_required('reports')
    def relatorio_gestao_excel():
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        report=management_data(request.args); workbook=Workbook(); sheet=workbook.active; sheet.title='Entradas e conclusões'
        sheet.append(['Período',report['start'].isoformat(),report['end'].isoformat()])
        sheet.append(['Data do registro','Entradas (data do pedido)','Conclusões (data registrada)'])
        for row in report['daily']: sheet.append([row['data'].isoformat(),row['entradas'],row['conclusoes']])
        sheet=workbook.create_sheet('Tempos por setor')
        sheet.append(['Setor','Encaminhamentos','Recebidos','Ainda aguardando','Encerrados sem recebimento',
            'Espera média (h)','Base espera','Atendimento médio encerrado (h)','Base atendimento','Sem marco final'])
        for row in report['sectors']:
            name=row['nome']; name="'"+name if name.startswith(('=','+','-','@')) else name
            sheet.append([name,row['enviados'],row['recebidos'],row['pendentes'],row['redirecionados'],row['espera_media'],
                row['base_espera'],row['atendimento_medio'],row['base_atendimento'],row['sem_marco']])
        notes=workbook.create_sheet('Critérios')
        for note in ('Horário de Brasília. Espera: envio até recebimento. Atendimento: recebimento até saída ou encerramento.',
                     'Médias setoriais consideram somente intervalos completos, sem durações negativas.',
                     'Cada encaminhamento conta separadamente, inclusive participações paralelas do mesmo processo.',
                     'Encaminhamentos selecionados pela data de envio; pendências são a situação atual desses encaminhamentos.',
                     'Conclusões usam a data registrada no sistema; não inferem o dia histórico a partir da situação.',
                     f"Encerrados sem data de conclusão em todo o acervo acessível: {report['sem_data_conclusao']}.",
                     f"Média até a conclusão: {report['media_dias']} dias corridos; base: {report['base_media']} processos."):
            notes.append([note])
        for sheet in workbook:
            first=2 if sheet.title=='Entradas e conclusões' else 1
            sheet.freeze_panes=f'A{first+1}'
            if sheet.title!='Critérios': sheet.auto_filter.ref=f'A{first}:{sheet.cell(sheet.max_row,sheet.max_column).coordinate}'
            for cell in sheet[first]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='176D63')
            for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col)+2))
        output=io.BytesIO(); workbook.save(output); output.seek(0)
        response=send_file(output,as_attachment=True,download_name=f"gestao-{report['start']}-{report['end']}.xlsx",
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response.headers['Cache-Control']='private, no-store'
        return response
