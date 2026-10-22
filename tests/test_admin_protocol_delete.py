"""Exclusão administrativa auditada, com dados inteiramente fictícios."""
from datetime import date, datetime, timezone
import hashlib
from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import load_workbook
import test_multitenancy as base
from app import app, db, gerar_proximo_numero_protocolo, verify_audit_chain
from models import (Anexo, DespachoProtocolo, EmissaoEletronica, HistoricoProtocolo,
                    Lotacao, Movimentacao, Organizacao, Protocolo, Usuario)


class AdminProtocolDeleteTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
        Path(base._test_directory.name).mkdir(exist_ok=True)
        with app.app_context():
            assert db.engine.dialect.name == 'sqlite'
            db.drop_all()
            db.create_all()
            db.session.add_all([Organizacao(id=1, nome='A', slug='a'),
                                Organizacao(id=2, nome='B', slug='b')])
            sector = Lotacao(id=1, tenant_id=1, nome='Setor fictício')
            db.session.add(sector)
            self.users, self.user_ids = {}, {}
            for role in ('admin', 'protocolista', 'tramitador', 'consulta'):
                user = Usuario(tenant_id=1, nome=role, login=role, tipo=role,
                               senha='fictitious-test-only', lotacao_id=1)
                db.session.add(user)
                db.session.flush()
                self.users[role], self.user_ids[role] = user.get_id(), user.id
            global_admin = Usuario(nome='Global', login='global', tipo='admin',
                                   senha='fictitious-test-only', is_platform_admin=True)
            db.session.add(global_admin)
            db.session.flush()
            self.users['global'] = global_admin.get_id()
            self.year = date.today().year
            protocol = Protocolo(tenant_id=1, numero=f'0042/{self.year}',
                                 nome='Alvo fictício da exclusão', matricula='FICTICIO',
                                 data_solicitacao=date.today(), status='EM ANÁLISE',
                                 setor_atual_id=1, responsavel='admin')
            foreign = Protocolo(tenant_id=2, numero=f'0042/{self.year}',
                                nome='Protocolo de outro cliente')
            db.session.add_all([protocol, foreign])
            db.session.flush()
            self.pid, self.foreign, self.token = protocol.id, foreign.id, protocol.consulta_token
            data = b'%PDF-1.7\nfictitious-preserved-document'
            digest = hashlib.sha256(data).hexdigest()
            attachment = Anexo(tenant_id=1, protocolo_id=self.pid,
                               file_name='ficticio.pdf', storage_path='ficticio.pdf',
                               storage_backend='database', file_data=data, file_hash=digest,
                               file_size=len(data), mime_type='application/pdf')
            db.session.add(attachment)
            db.session.flush()
            self.aid, self.data = attachment.id, data
            db.session.add(DespachoProtocolo(
                tenant_id=1, protocolo_id=self.pid, numero=1,
                protocolo_numero=protocol.numero, texto='Despacho fictício',
                autor_id=self.user_ids['admin'], autor_nome='admin',
                criado_em=datetime.now(timezone.utc), envio_token='fictitious-dispatch-token',
                file_name='despacho.pdf', storage_path='despacho.pdf',
                storage_backend='database', file_hash=digest, file_data=data))
            db.session.add(HistoricoProtocolo(
                tenant_id=1, protocolo_id=self.pid, acao='CRIACAO',
                status=protocol.status, responsavel='admin', observacao='Registro preservado.'))
            db.session.commit()
        self.client = app.test_client()

    def tearDown(self):
        app.config['WTF_CSRF_ENABLED'] = False
        with app.app_context():
            db.session.remove()
            db.drop_all()

    def login(self, role='admin', tenant=None):
        with self.client.session_transaction() as session:
            session.clear()
            session['_user_id'], session['_fresh'] = self.users[role], True
            if tenant:
                session['active_tenant_id'] = tenant

    def delete(self, pid=None, **changes):
        data = {'justificativa': 'Protocolo duplicado criado para teste.', 'confirmar_exclusao': '1'}
        data.update(changes)
        return self.client.post(f'/protocolo/{pid or self.pid}/deletar', data=data)

    def test_admin_deletes_once_and_preserves_documents_and_audit(self):
        self.login()
        result = self.delete()
        self.assertEqual(result.status_code, 302)
        self.assertTrue(result.location.endswith('/protocolos'))
        with app.app_context():
            protocol = db.session.get(Protocolo, self.pid)
            self.assertIsNotNone(protocol.excluido_em)
            self.assertEqual(protocol.status, 'EM ANÁLISE')
            self.assertEqual(db.session.get(Anexo, self.aid).file_data, self.data)
            self.assertEqual(DespachoProtocolo.query.one().file_data, self.data)
            self.assertEqual(HistoricoProtocolo.query.filter_by(acao='CRIACAO').count(), 1)
            event = HistoricoProtocolo.query.filter_by(acao='EXCLUSAO').one()
            self.assertEqual(event.usuario_id, self.user_ids['admin'])
            self.assertIn('Protocolo duplicado', event.observacao)
            self.assertTrue(verify_audit_chain(1)[0])
            self.assertIsNone(db.session.get(Protocolo, self.foreign).excluido_em)
        self.assertEqual(self.delete().status_code, 404)

    def test_button_only_for_admin_even_when_archived_or_awaiting_receipt(self):
        self.login()
        for state in ('normal', 'archived', 'portal'):
            with app.app_context():
                protocol = db.session.get(Protocolo, self.pid)
                protocol.arquivado_em = datetime.now() if state == 'archived' else None
                protocol.modalidade_abertura = 'remota_requerente' if state == 'portal' else 'presencial_protocolista'
                db.session.commit()
            html = self.client.get(f'/protocolo/{self.pid}').get_data(as_text=True)
            self.assertIn('data-bs-target="#modalExcluirProtocolo"', html)
            self.assertIn('name="csrf_token"', html)
            self.assertIn('name="confirmar_exclusao"', html)
        for role in ('protocolista', 'tramitador', 'consulta'):
            self.login(role)
            html = self.client.get(f'/protocolo/{self.pid}').get_data(as_text=True)
            self.assertNotIn('modalExcluirProtocolo', html)
            self.assertIn(self.delete().status_code, (302, 403))
            with app.app_context():
                self.assertIsNone(db.session.get(Protocolo, self.pid).excluido_em)

    def test_confirmation_reason_post_and_csrf_are_required(self):
        self.login()
        self.assertEqual(self.client.get(f'/protocolo/{self.pid}/deletar').status_code, 405)
        for data in ({'confirmar_exclusao': ''}, {'justificativa': ''},
                     {'justificativa': 'curto'}, {'justificativa': 'x' * 2001},
                     {'justificativa': 'motivo com nulo\x00'}):
            self.assertEqual(self.delete(**data).status_code, 400)
        app.config['WTF_CSRF_ENABLED'] = True
        self.assertEqual(self.delete().status_code, 400)
        with app.app_context():
            self.assertIsNone(db.session.get(Protocolo, self.pid).excluido_em)
            self.assertEqual(HistoricoProtocolo.query.filter_by(acao='EXCLUSAO').count(), 0)

    def test_other_client_anonymous_and_unselected_global_cannot_delete(self):
        self.assertEqual(self.delete().status_code, 302)
        self.login()
        self.assertEqual(self.delete(pid=self.foreign).status_code, 404)
        self.login('global')
        self.assertEqual(self.delete().status_code, 302)
        with app.app_context():
            self.assertIsNone(db.session.get(Protocolo, self.pid).excluido_em)
        self.login('global', tenant=1)
        self.assertEqual(self.delete().status_code, 302)
        with app.app_context():
            self.assertIsNotNone(db.session.get(Protocolo, self.pid).excluido_em)

    def test_deleted_protocol_hidden_from_queries_files_and_reports(self):
        self.login()
        self.delete()
        for path in (f'/protocolo/{self.pid}', f'/protocolo/{self.pid}/pdf',
                     f'/api/protocolo/{self.pid}', f'/anexo/{self.aid}/download',
                     f'/protocolo/{self.pid}/despachos/1/pdf', f'/consulta/{self.token}'):
            self.assertEqual(self.client.get(path).status_code, 404, path)
        for action in ('editar', 'anexo/novo', 'tramitar', 'receber', 'desarquivar', 'despachos'):
            path = f'/protocolo/{self.pid}/{action}'
            self.assertEqual(self.client.post(path, data={}).status_code, 404, path)
        for path in ('/protocolos', '/relatorios'):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200)
            self.assertNotIn('Alvo fictício da exclusão', response.get_data(as_text=True))
        stats = self.client.get('/protocolos/dashboard-stats')
        self.assertEqual(stats.status_code, 200)
        self.assertEqual(stats.json['statusProtocolos'], [])
        report = self.client.get('/protocolos/backup/excel')
        book = load_workbook(BytesIO(report.data))
        self.assertEqual(book.active.max_row, 1)
        audit = self.client.get('/admin/auditoria?exportar=1')
        self.assertEqual(audit.status_code, 200)
        self.assertTrue(audit.json['integra'])
        self.assertTrue(any(e['payload']['acao'] == 'EXCLUSAO' for e in audit.json['eventos']))

    def test_deleted_number_not_reused_and_pending_movement_preserved(self):
        self.login()
        with app.app_context():
            db.session.add(Movimentacao(tenant_id=1, protocolo_id=self.pid,
                setor_destino_id=1, enviado_por_id=self.user_ids['admin']))
            db.session.commit()
        self.delete()
        self.assertNotIn(f'0042/{self.year}', self.client.get('/pendencias-recebimento').get_data(as_text=True))
        self.assertEqual(self.client.get(f'/protocolos/ultimoNumero/{self.year}').json, {'ultimo': 42})
        with self.client:
            self.client.get('/protocolos')
            self.assertEqual(gerar_proximo_numero_protocolo(), f'0043/{self.year}')
        with app.app_context():
            self.assertEqual(Movimentacao.query.filter_by(protocolo_id=self.pid).count(), 1)

    def test_public_validation_hidden_without_rewriting_signed_document(self):
        with app.app_context():
            db.session.add(EmissaoEletronica(
                tenant_id=1, protocolo_id=self.pid, emitido_por_id=self.user_ids['admin'],
                pdf_anexo_id=self.aid, token_publico='fictitious-public-validation',
                codigo_publico='FICTICIO', pdf_sha256=hashlib.sha256(self.data).hexdigest(),
                declaracao='Teste fictício', nome_emitente='admin', login_emitente='admin',
                ip_hash='0' * 64, user_agent_hash='0' * 64))
            db.session.commit()
        self.login()
        self.assertEqual(self.client.get('/validar-emissao/fictitious-public-validation').status_code, 200)
        self.delete()
        self.assertEqual(self.client.get('/validar-emissao/fictitious-public-validation').status_code, 404)
        with app.app_context():
            self.assertEqual(EmissaoEletronica.query.one().status, 'VALIDA')
            self.assertEqual(db.session.get(Anexo, self.aid).file_data, self.data)
