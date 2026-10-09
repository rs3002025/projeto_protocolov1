from pathlib import Path
from datetime import datetime, timezone, timedelta
from types import SimpleNamespace
import io
from openpyxl import load_workbook
import test_multitenancy as base
from app import app, db, format_brasilia, chronological
from models import Servidor, Protocolo, HistoricoProtocolo

def setup_module():
    Path(base._test_directory.name).mkdir(exist_ok=True)
    base.setup_module()

def teardown_module():
    base.teardown_module()

def test_brasilia_converts_utc_and_preserves_aware_offset():
    assert format_brasilia(datetime(2026,10,8,1,30)) == '07/10/2026 22:30'
    assert format_brasilia(datetime(2026,10,8,11,tzinfo=timezone.utc)) == '08/10/2026 08:00'
    assert format_brasilia(datetime(2026,10,8,8,tzinfo=timezone(timedelta(hours=-3)))) == '08/10/2026 08:00'
    assert format_brasilia(None) == ''

def test_timeline_oldest_first_and_tie_by_id():
    events=[SimpleNamespace(id=i,data_movimentacao=dt) for i,dt in
            [(3,datetime(2026,10,8,12)),(2,datetime(2026,10,8,11)),(1,datetime(2026,10,8,11))]]
    assert [e.id for e in chronological(events)] == [1,2,3]
    for name in ['protocolo_detalhe.html','portal_detalhe.html']:
        source=Path(app.root_path,'templates',name).read_text(encoding='utf-8')
        assert 'protocolo.historico|chronological' in source

def test_download_fill_import_and_prefill_all_fields_is_tenant_scoped():
    client=app.test_client(); base.login(client,'cliente-a')
    response=client.get('/admin/servidores/modelo')
    assert response.status_code == 200
    book=load_workbook(io.BytesIO(response.data))
    sheet=book['Servidores']
    headers=[c.value for c in sheet[1]]
    record={'matricula':'00042','nome':'Pessoa Importada','cpf':'01234567890',
            'data_nascimento':datetime(1990,1,1),'nome_mae':'Maria Importada',
            'rg':'00123','cargo':'Analista','lotacao':'SEAD','unidade_de_exercicio':'Unidade Teste',
            'endereco':'Rua Teste, 10','bairro':'Centro','municipio':'Morada Nova/CE',
            'cep':'01234000','telefone':'88999999999','email':'pessoa@example.test'}
    assert set(record) == set(headers)
    for i,header in enumerate(headers,1): sheet.cell(2,i,record[header])
    assert sheet['A2'].number_format == '@'
    book.active=1  # Importador usa a aba certa mesmo com Orientações selecionada.
    stream=io.BytesIO(); book.save(stream); stream.seek(0)
    result=client.post('/admin/servidores/importar',data={'arquivo':(stream,'servidores.xlsx')},follow_redirects=True)
    assert result.status_code == 200
    with app.app_context():
        server=Servidor.query.filter_by(tenant_id=1,matricula='00042').one()
        for field in ['rg','cargo','lotacao','unidade_de_exercicio','endereco','bairro','municipio','cep','telefone','email']:
            assert getattr(server,field) == record[field]
    data=client.get('/api/servidor/00042').get_json()
    for field in ['cpf','rg','endereco','bairro','municipio','cep','telefone']:
        assert data[field] == record[field]
    other=app.test_client(); base.login(other,'cliente-b')
    assert other.get('/api/servidor/00042').status_code == 404
    assert app.test_client().get('/admin/servidores/modelo').status_code == 302

def test_public_history_is_chronological_and_brasilia():
    with app.app_context():
        protocol=Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        token=protocol.consulta_token
        db.session.add_all([HistoricoProtocolo(tenant_id=1,protocolo_id=protocol.id,
            acao='EVENTO_POSTERIOR',status='EM ANÁLISE',responsavel='admin',data_movimentacao=datetime(2026,10,8,12)),
            HistoricoProtocolo(tenant_id=1,protocolo_id=protocol.id,acao='EVENTO_ANTERIOR',
            status='RECEBIDO',responsavel='admin',data_movimentacao=datetime(2026,10,8,11))])
        db.session.commit()
    html=app.test_client().post(f'/consulta/{token}',data={'matricula':'MAT-A'}).get_data(as_text=True)
    assert html.index('Evento Anterior') < html.index('Evento Posterior')
    assert '08/10/2026 08:00' in html and '08/10/2026 09:00' in html

def test_public_history_uses_reader_labels_and_pending_receipt_state():
    with app.app_context():
        protocol=Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        token=protocol.consulta_token
        previous=protocol.modalidade_abertura
        protocol.modalidade_abertura='remota_requerente'
        db.session.add(HistoricoProtocolo(tenant_id=1,protocolo_id=protocol.id,
            acao='ENVIO_REMOTO',status='AGUARDANDO RECEBIMENTO',responsavel='servidor'))
        db.session.commit()
    try:
        html=app.test_client().post(f'/consulta/{token}',data={'matricula':'MAT-A'}).get_data(as_text=True)
        assert 'Requerimento enviado pelo portal' in html
        assert 'ENVIO_REMOTO' not in html
        assert '<dd>AGUARDANDO RECEBIMENTO</dd>' in html
    finally:
        with app.app_context():
            protocol=Protocolo.query.filter_by(nome='Dado exclusivo A').one()
            protocol.modalidade_abertura=previous
            db.session.commit()
