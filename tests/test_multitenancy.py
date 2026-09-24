import os
import tempfile
import io
import sys
import types
import hashlib
from unittest.mock import patch
from openpyxl import load_workbook
from PIL import Image
from flask import render_template
from pathlib import Path
from datetime import date, datetime, timedelta

# Nunca herdar DATABASE_URL do Railway: esta suíte recria todas as tabelas.
_test_directory = tempfile.TemporaryDirectory(prefix='protocolo-tests-')
os.environ['SECRET_KEY'] = 'test-secret-key'
os.environ['DATABASE_URL'] = 'sqlite:///' + str(Path(_test_directory.name) / 'tests.sqlite3')

from app import app, bcrypt, db
from models import (Lotacao, Movimentacao, Organizacao, Protocolo, Usuario,
                    ConsultaPublicaTentativa, LoginTentativa, Anexo,
                    HistoricoProtocolo, Servidor, TipoRequerimento, ChamadoSuporte, MensagemSuporte)
from models import OrganizacaoCapacidadeEvento
from models import EmissaoEletronica


def teardown_module():
    with app.app_context():
        db.session.remove()
        db.engine.dispose()
    _test_directory.cleanup()


def setup_module():
    app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    with app.app_context():
        assert db.engine.url.drivername == 'sqlite'
        assert Path(db.engine.url.database).resolve().parent == Path(_test_directory.name).resolve()
        db.drop_all()
        db.create_all()
        a = Organizacao(nome='Cliente A', slug='cliente-a')
        b = Organizacao(nome='Cliente B', slug='cliente-b')
        db.session.add_all([a, b])
        db.session.flush()
        password = bcrypt.generate_password_hash('senha-segura').decode('utf-8')
        db.session.add_all([
            Usuario(tenant_id=a.id, nome='Ana', nome_completo='Ana A', login='admin', senha=password, tipo='admin'),
            Usuario(tenant_id=b.id, nome='Bia', nome_completo='Bia B', login='admin', senha=password, tipo='admin'),
            Usuario(tenant_id=a.id, nome='Cris', nome_completo='Cris Consulta', login='consulta', senha=password, tipo='consulta'),
        ])
        db.session.add_all([
            Lotacao(tenant_id=a.id, nome='Protocolo'),
            Lotacao(tenant_id=a.id, nome='Jurídico'),
        ])
        db.session.flush()
        db.session.add_all([
            Protocolo(tenant_id=a.id, numero='0001/2026', nome='Dado exclusivo A', matricula='MAT-A', data_solicitacao=date.today()),
            Protocolo(tenant_id=b.id, numero='0001/2026', nome='Dado exclusivo B', matricula='MAT-B', data_solicitacao=date.today()),
        ])
        db.session.commit()


def login(client, organizacao):
    return client.post('/login', data={
        'organizacao': organizacao,
        'login': 'admin',
        'senha': 'senha-segura',
    }, follow_redirects=True)


def test_login_por_cliente_nao_exibe_escolha_de_organizacao():
    client = app.test_client()
    pagina = client.get('/entrar/cliente-a')
    html = pagina.get_data(as_text=True)
    assert pagina.status_code == 200
    assert 'Acesso de Cliente A' in html
    assert 'name="organizacao"' not in html

    response = client.post('/entrar/cliente-a', data={
        'login': 'admin', 'senha': 'senha-segura',
    })
    assert response.status_code == 302
    assert response.headers['Location'] == '/home'


def test_administrador_geral_controla_capacidades_com_auditoria():
    with app.app_context():
        admin = Usuario.query.filter_by(tenant_id=1, login='admin').one()
        admin.is_platform_admin = True
        db.session.commit()

    client = app.test_client()
    client.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura',
    })
    response = client.post('/plataforma/cliente/2/capacidades', data={
        'emissao_eletronica_protocolista_enabled': 'on',
        'portal_servidor_remoto_enabled': 'on',
        'nivel_garantia_assinatura': 'forte',
        'motivo': 'Homologação do recurso contratado',
    }, follow_redirects=True)
    assert response.status_code == 200
    assert 'Capacidades de Cliente B atualizadas e auditadas.' in response.get_data(as_text=True)

    with app.app_context():
        cliente = Organizacao.query.filter_by(slug='cliente-b').one()
        assert cliente.emissao_eletronica_protocolista_enabled is True
        assert cliente.portal_servidor_remoto_enabled is True
        assert cliente.nivel_garantia_assinatura == 'forte'
        evento = OrganizacaoCapacidadeEvento.query.filter_by(organizacao_id=cliente.id).one()
        assert evento.alterado_por_id == 1
        assert evento.motivo == 'Homologação do recurso contratado'

        OrganizacaoCapacidadeEvento.query.delete()
        cliente.emissao_eletronica_protocolista_enabled = False
        cliente.portal_servidor_remoto_enabled = False
        cliente.nivel_garantia_assinatura = 'interno'
        Usuario.query.filter_by(tenant_id=1, login='admin').one().is_platform_admin = False
        db.session.commit()


def test_login_nao_redireciona_para_site_externo():
    client = app.test_client()
    response = client.post('/login?next=https://evil.example/coleta', data={
        'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura',
    })
    assert response.status_code == 302
    assert response.headers['Location'] == '/home'

    client.post('/logout')
    response = client.post('/login?next=/protocolos', data={
        'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura',
    })
    assert response.headers['Location'] == '/protocolos'


def test_login_bloqueia_forca_bruta_sem_revelar_usuario():
    client = app.test_client()
    dados = {'organizacao': 'cliente-a', 'login': 'inexistente', 'senha': 'incorreta'}
    for _ in range(5):
        response = client.post('/login', data=dados)
        assert response.status_code == 200
        assert 'Não foi possível autenticar' in response.get_data(as_text=True)
    bloqueado = client.post('/login', data=dados)
    assert bloqueado.status_code == 429
    assert 'Aguarde alguns minutos' in bloqueado.get_data(as_text=True)
    with app.app_context():
        tentativa = LoginTentativa.query.one()
        tentativa.janela_iniciada_em = datetime.utcnow() - timedelta(minutes=16)
        tentativa.bloqueado_ate = None
        db.session.commit()
    assert client.post('/login', data=dados).status_code == 200


def test_health_verifica_a_conexao_com_o_banco():
    client = app.test_client()
    response = client.get('/health')
    assert response.status_code == 200
    assert response.get_json() == {'status': 'ok', 'database': 'ok'}
    with patch.object(db.session, 'execute', side_effect=RuntimeError('banco indisponivel')):
        response = client.get('/health')
    assert response.status_code == 503
    assert response.get_json() == {'status': 'indisponivel', 'database': 'erro'}


def test_cabecalhos_protegem_dados_e_transporte():
    client = app.test_client()
    login(client, 'cliente-a')
    response = client.get('/protocolos', base_url='https://localhost')
    assert response.headers['Cache-Control'] == 'no-store, max-age=0'
    assert response.headers['Pragma'] == 'no-cache'
    assert response.headers['X-Content-Type-Options'] == 'nosniff'
    assert response.headers['X-Frame-Options'] == 'DENY'
    assert response.headers['Strict-Transport-Security'].startswith('max-age=31536000')
    assert "object-src 'none'" in response.headers['Content-Security-Policy']
    assert "connect-src 'self' https://viacep.com.br" in response.headers['Content-Security-Policy']
    assert response.headers['Permissions-Policy'] == 'camera=(), microphone=(), geolocation=()'
    assert app.config['MAX_CONTENT_LENGTH'] == 21 * 1024 * 1024


def test_listagem_nao_vaza_dados_entre_clientes():
    client = app.test_client()
    login(client, 'cliente-a')
    response = client.get('/protocolos')
    assert response.status_code == 200
    assert b'Dado exclusivo A' in response.data
    assert b'Dado exclusivo B' not in response.data


def test_acesso_direto_a_protocolo_de_outro_cliente_retorna_404():
    with app.app_context():
        protocolo_b = Protocolo.query.filter_by(nome='Dado exclusivo B').one()
        protocolo_b_id = protocolo_b.id
    client = app.test_client()
    login(client, 'cliente-a')
    assert client.get(f'/protocolo/{protocolo_b_id}').status_code == 404


def test_todas_as_operacoes_de_protocolo_bloqueiam_id_de_outro_cliente():
    with app.app_context():
        protocolo_b_id = Protocolo.query.filter_by(nome='Dado exclusivo B').one().id
        setor_a_id = Lotacao.query.filter_by(tenant_id=1, nome='Jurídico').one().id
    client = app.test_client()
    login(client, 'cliente-a')

    requisicoes = [
        ('get', f'/protocolo/{protocolo_b_id}'),
        ('get', f'/api/protocolo/{protocolo_b_id}'),
        ('get', f'/protocolo/{protocolo_b_id}/pdf'),
        ('post', f'/protocolo/{protocolo_b_id}/documento/gerar'),
        ('post', f'/protocolo/{protocolo_b_id}/editar'),
        ('post', f'/protocolo/{protocolo_b_id}/deletar'),
        ('post', f'/protocolo/{protocolo_b_id}/anexo/novo'),
        ('post', f'/protocolo/{protocolo_b_id}/tramitar'),
        ('post', f'/protocolo/{protocolo_b_id}/receber'),
        ('post', f'/protocolo/{protocolo_b_id}/arquivar'),
    ]
    for metodo, caminho in requisicoes:
        dados = {'setor_destino_id': setor_a_id} if caminho.endswith('/tramitar') else {}
        response = getattr(client, metodo)(caminho, data=dados)
        assert response.status_code == 404, caminho
    response = client.post('/protocolos/atualizar', json={
        'protocoloId': protocolo_b_id, 'novoStatus': 'EM ANÁLISE'
    })
    assert response.status_code == 404

    with app.app_context():
        protocolo_b = db.session.get(Protocolo, protocolo_b_id)
        assert protocolo_b.nome == 'Dado exclusivo B'
        assert protocolo_b.status == 'Aberto'
        assert protocolo_b.arquivado_em is None


def test_protocolo_legado_sem_data_continua_consultavel():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        organizacao = Organizacao.query.filter_by(slug='cliente-a').one()
        legado = Protocolo(
            tenant_id=organizacao.id, numero='0999/2024', nome='Registro legado sem data',
            matricula='LEGADO-SEM-DATA', data_solicitacao=None,
        )
        db.session.add(legado)
        db.session.commit()
        protocolo_id = legado.id
    for rota in (f'/protocolo/{protocolo_id}', '/protocolos', '/relatorios'):
        response = client.get(rota)
        assert response.status_code == 200
        assert 'Não informada' in response.get_data(as_text=True)
        assert 'Não definido' in response.get_data(as_text=True)

def test_apis_de_cadastro_e_bairros_nao_vazam_dados_de_outro_cliente():
    with app.app_context():
        protocolo_a = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo_b = Protocolo.query.filter_by(nome='Dado exclusivo B').one()
        protocolo_a.bairro = 'Bairro exclusivo A'
        protocolo_b.bairro = 'Bairro secreto B'
        db.session.add(Servidor(tenant_id=protocolo_b.tenant_id, matricula='SERV-B',
                                nome='Servidor exclusivo B'))
        db.session.commit()

    client = app.test_client()
    login(client, 'cliente-a')
    bairros = client.get('/api/bairros').get_json()
    assert 'Bairro exclusivo A' in bairros
    assert 'Bairro secreto B' not in bairros
    assert 'Outro' in bairros
    assert client.get('/api/servidor/SERV-B').status_code == 404
    assert client.get('/api/servidores/search?nome=Servidor').get_json() == []
    usuarios = client.get('/api/usuarios').get_json()
    assert all(usuario['login'] != 'admin' or usuario['nome'] == 'Ana' for usuario in usuarios)

    formulario = client.get('/protocolo/novo').get_data(as_text=True)
    assert 'value="Morada Nova"' not in formulario
    assert 'list="bairrosDisponiveis"' in formulario


def test_mesmo_login_e_numero_podem_existir_em_clientes_distintos():
    with app.app_context():
        assert Usuario.query.filter_by(login='admin').count() == 2
        assert Protocolo.query.filter_by(numero='0001/2026').count() == 2


def test_consulta_publica_usa_organizacao_e_nao_expoe_requerente():
    with app.app_context():
        token = Protocolo.query.filter_by(nome='Dado exclusivo A').one().consulta_token
    client = app.test_client()
    response = client.get(f'/consulta/{token}')
    assert response.status_code == 200
    assert b'0001/2026' not in response.data
    assert b'Dado exclusivo A' not in response.data
    response = client.post(f'/consulta/{token}', data={'matricula': 'incorreta'})
    assert b'0001/2026' not in response.data
    response = client.post(f'/consulta/{token}', data={'matricula': 'MAT-A'})
    assert b'0001/2026' in response.data
    assert b'Dado exclusivo A' not in response.data
    assert client.get('/consulta/token-inexistente').status_code == 404
    assert client.get('/consulta/cliente-a/2026/0001').status_code == 404


def test_localizacao_informa_tramite_pendente_sem_transferir_custodia():
    with app.app_context():
        tenant = Organizacao.query.filter_by(slug='cliente-a').one()
        origem = Lotacao.query.filter_by(tenant_id=tenant.id, nome='Protocolo').one()
        destino = Lotacao.query.filter_by(tenant_id=tenant.id, nome='Jurídico').one()
        admin = Usuario.query.filter_by(tenant_id=tenant.id, login='admin').one()
        protocolo = Protocolo(
            tenant_id=tenant.id, numero='LOCAL-1/2026', nome='Teste de localização',
            matricula='LOCAL-SEGURA', data_solicitacao=date.today(),
            setor_atual_id=origem.id, status='EM TRAMITAÇÃO')
        db.session.add(protocolo)
        db.session.flush()
        db.session.add(Movimentacao(
            tenant_id=tenant.id, protocolo_id=protocolo.id,
            setor_origem_id=origem.id, setor_destino_id=destino.id,
            enviado_por_id=admin.id))
        db.session.commit()
        protocolo_id, token = protocolo.id, protocolo.consulta_token

    client = app.test_client()
    login(client, 'cliente-a')
    detalhe = client.get(f'/protocolo/{protocolo_id}').get_data(as_text=True)
    listagem = client.get('/protocolos?numero=LOCAL-1').get_data(as_text=True)
    relatorio = client.get('/relatorios?numero=LOCAL-1').get_data(as_text=True)
    api = client.get(f'/api/protocolo/{protocolo_id}').get_json()
    assert 'Em trânsito de Protocolo para Jurídico (aguardando recebimento)' in detalhe
    assert 'Em trânsito de Protocolo para Jurídico (aguardando recebimento)' in listagem
    assert 'Em trânsito de Protocolo para Jurídico (aguardando recebimento)' in relatorio
    assert api['localizacao_atual'] == 'Em trânsito de Protocolo para Jurídico (aguardando recebimento)'

    publico = app.test_client().post(
        f'/consulta/{token}', data={'matricula': 'LOCAL-SEGURA'}).get_data(as_text=True)
    assert 'Em trânsito de Protocolo para Jurídico (aguardando recebimento)' in publico

    with app.app_context():
        protocolo = db.session.get(Protocolo, protocolo_id)
        assert protocolo.setor_atual.nome == 'Protocolo'
        db.session.delete(protocolo)
        db.session.commit()


def test_consulta_publica_bloqueia_forca_bruta():
    with app.app_context():
        token = Protocolo.query.filter_by(nome='Dado exclusivo B').one().consulta_token
    client = app.test_client()
    for _ in range(5):
        assert client.post(f'/consulta/{token}', data={'matricula': 'incorreta'}).status_code == 200
    assert client.post(f'/consulta/{token}', data={'matricula': 'incorreta'}).status_code == 429


def test_consulta_publica_reutiliza_janela_expirada_e_aceita_unicode():
    client = app.test_client()
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        token, protocolo_id = protocolo.consulta_token, protocolo.id
    assert client.post(f'/consulta/{token}', data={'matricula': 'inválida'}).status_code == 200
    with app.app_context():
        tentativa = ConsultaPublicaTentativa.query.filter_by(protocolo_id=protocolo_id).one()
        tentativa.janela_iniciada_em = datetime.utcnow() - timedelta(minutes=16)
        db.session.commit()
    assert client.post(f'/consulta/{token}', data={'matricula': 'outra'}).status_code == 200
    with app.app_context():
        tentativa = ConsultaPublicaTentativa.query.filter_by(protocolo_id=protocolo_id).one()
        assert tentativa.tentativas == 1
    assert b'0001/2026' in client.post(f'/consulta/{token}', data={'matricula': 'MAT-A'}).data


def test_relatorios_e_exportacao_exigem_permissao():
    client = app.test_client()
    client.post('/login', data={'organizacao': 'cliente-a', 'login': 'consulta', 'senha': 'senha-segura'})
    for path in ['/relatorios', '/protocolos/backup/excel']:
        assert client.get(path, headers={'Content-Type': 'application/json'}).status_code == 403
    for path in ['/api/usuarios', '/api/servidor/MAT-A', '/api/servidores/search?nome=Ana',
                 '/protocolos/ultimoNumero/2026']:
        assert client.get(path, headers={'Content-Type': 'application/json'}).status_code == 403
    html = client.get('/protocolos').get_data(as_text=True)
    assert 'Exportar para Excel' not in html
    assert '>Relatórios</a>' not in html
    login(client, 'cliente-a')  # A sessão de consulta não deve ser elevada pelo formulário.
    assert client.get('/relatorios', headers={'Content-Type': 'application/json'}).status_code == 403


def test_logout_exige_post_e_csrf_quando_habilitado():
    client = app.test_client()
    login(client, 'cliente-a')
    assert client.get('/logout').status_code == 405
    app.config['WTF_CSRF_ENABLED'] = True
    try:
        assert client.post('/logout').status_code == 400
        html = client.get('/home').get_data(as_text=True)
        import re
        token = re.search(r'action="/logout"[\s\S]*?name="csrf_token" value="([^"]+)"', html).group(1)
        response = client.post('/logout', data={'csrf_token': token})
        assert response.status_code == 302
        assert client.get('/home').status_code == 302
    finally:
        app.config['WTF_CSRF_ENABLED'] = False


def test_csrf_rejeita_operacao_sem_token_e_aceita_formulario_legitimo():
    import re
    client = app.test_client()
    app.config['WTF_CSRF_ENABLED'] = True
    try:
        assert client.post('/login', data={}).status_code == 400
        html = client.get('/login').get_data(as_text=True)
        token = re.search(r'name="csrf_token"[^>]*value="([^"]+)"', html).group(1)
        response = client.post('/login', data={
            'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura',
            'csrf_token': token,
        })
        assert response.status_code == 302
        assert client.post('/protocolos/atualizar', json={}).status_code == 400
        response = client.post('/protocolos/atualizar', json={}, headers={'X-CSRFToken': token})
        assert b'CSRF' not in response.data
        with app.app_context():
            protocolo_id = Protocolo.query.filter_by(nome='Dado exclusivo A').one().id
        for path in ['/protocolo/novo', '/configuracoes', f'/protocolo/{protocolo_id}']:
            response = client.get(path)
            assert response.status_code == 200
            for form in re.findall(r'<form\b[\s\S]*?</form>', response.get_data(as_text=True)):
                if re.search(r'method="POST"', form, re.I):
                    assert 'name="csrf_token"' in form
    finally:
        app.config['WTF_CSRF_ENABLED'] = False


def test_login_https_preserva_referer_necessario_ao_csrf():
    import re
    client = app.test_client()
    app.config['WTF_CSRF_ENABLED'] = True
    try:
        pagina = client.get('/login', base_url='https://sysprot.example')
        assert pagina.headers['Referrer-Policy'] == 'same-origin'
        token = re.search(
            r'name="csrf_token"[^>]*value="([^"]+)"', pagina.get_data(as_text=True)
        ).group(1)
        response = client.post(
            '/login',
            base_url='https://sysprot.example',
            headers={'Referer': 'https://sysprot.example/login'},
            data={
                'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura',
                'csrf_token': token,
            },
        )
        assert response.status_code == 302
    finally:
        app.config['WTF_CSRF_ENABLED'] = False


def test_logo_personalizada_fica_isolada_por_organizacao():
    client = app.test_client()
    login(client, 'cliente-a')
    primeira = Path('static/img/logo.png').read_bytes()
    segunda = Path('static/img/logobrasao.png').read_bytes()
    response = client.post('/admin/identidade/logo', data={
        'logo': (io.BytesIO(primeira), 'logo.png'),
        'salvar': 'Salvar logo',
    }, content_type='multipart/form-data', follow_redirects=True)
    assert response.status_code == 200
    with app.app_context():
        primeira_gravada = Organizacao.query.filter_by(slug='cliente-a').one().logo_data
        assert primeira_gravada
        assert Organizacao.query.filter_by(slug='cliente-b').one().logo_data is None
    response = client.post('/admin/identidade/logo', data={
        'logo': (io.BytesIO(segunda), 'brasao.png'),
        'salvar': 'Salvar logo',
    }, content_type='multipart/form-data', follow_redirects=True)
    assert response.status_code == 200
    with app.app_context():
        segunda_gravada = Organizacao.query.filter_by(slug='cliente-a').one().logo_data
        assert segunda_gravada and segunda_gravada != primeira_gravada
    response = client.get('/identidade/cliente-a/logo')
    assert response.status_code == 200
    assert response.content_type.startswith('image/png')
    response = client.post('/admin/identidade/logo', data={
        'remover': 'Restaurar logo padrão',
    }, content_type='multipart/form-data', follow_redirects=True)
    assert response.status_code == 200
    with app.app_context():
        assert Organizacao.query.filter_by(slug='cliente-a').one().logo_data is None
    response = client.get('/identidade/cliente-a/logo')
    assert response.status_code == 302
    assert 'logo-sysprot.svg' in response.headers['Location']


def test_logo_remove_margens_brancas_sem_distorcer_o_conteudo():
    client = app.test_client()
    login(client, 'cliente-a')
    source = io.BytesIO()
    image = Image.new('RGB', (1200, 600), 'white')
    for x in range(450, 750):
        for y in range(240, 360):
            image.putpixel((x, y), (0, 80, 90))
    image.save(source, format='PNG')
    source.seek(0)
    response = client.post('/admin/identidade/logo', data={
        'logo': (source, 'logo-com-margens.png'), 'salvar': 'Salvar logo',
    }, content_type='multipart/form-data', follow_redirects=True)
    assert response.status_code == 200
    logo = Image.open(io.BytesIO(client.get('/identidade/cliente-a/logo').data))
    assert logo.width < 400 and logo.height < 250


def test_tramitacao_registra_destino_e_historico():
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        destino = Lotacao.query.filter_by(nome='Jurídico').one()
        protocolo_id, destino_id = protocolo.id, destino.id
    client = app.test_client()
    login(client, 'cliente-a')
    response = client.post(f'/protocolo/{protocolo_id}/tramitar', data={
        'setor_destino_id': destino_id,
        'observacao': 'Análise jurídica necessária',
    })
    assert response.status_code == 302
    with app.app_context():
        movimento = Movimentacao.query.filter_by(protocolo_id=protocolo_id).one()
        assert movimento.setor_destino_id == destino_id
        assert movimento.recebido_em is None
        assert movimento.protocolo.status == 'EM TRAMITAÇÃO'


def test_tramitacao_pode_destinar_usuario_ou_todo_setor():
    with app.app_context():
        org = Organizacao.query.filter_by(slug='cliente-a').one()
        setor = Lotacao.query.filter_by(tenant_id=org.id, nome='Jurídico').one()
        senha = bcrypt.generate_password_hash('senha-segura').decode('utf-8')
        destinatario = Usuario(tenant_id=org.id, nome='João', login='joao', senha=senha,
                               tipo='tramitador', lotacao_id=setor.id)
        colega = Usuario(tenant_id=org.id, nome='Maria', login='maria', senha=senha,
                         tipo='tramitador', lotacao_id=setor.id)
        db.session.add_all([destinatario, colega])
        db.session.flush()
        protocolo = Protocolo(tenant_id=org.id, numero='DEST-1/2026', nome='Destino individual',
                              data_solicitacao=date.today())
        db.session.add(protocolo)
        db.session.commit()
        protocolo_id, setor_id, destinatario_id = protocolo.id, setor.id, destinatario.id
    admin = app.test_client()
    login(admin, 'cliente-a')
    resposta = admin.post(f'/protocolo/{protocolo_id}/tramitar', data={
        'setor_destino_id': setor_id, 'destinatario_usuario_id': destinatario_id})
    assert resposta.status_code == 302
    with app.app_context():
        movimento = Movimentacao.query.filter_by(protocolo_id=protocolo_id).one()
        assert movimento.destinatario_usuario_id == destinatario_id

    maria = app.test_client()
    maria.post('/login', data={'organizacao': 'cliente-a', 'login': 'maria', 'senha': 'senha-segura'})
    assert b'DEST-1/2026' not in maria.get('/pendencias-recebimento').data
    assert maria.post(f'/protocolo/{protocolo_id}/receber').status_code == 302
    with app.app_context():
        assert Movimentacao.query.filter_by(protocolo_id=protocolo_id).one().recebido_em is None

    joao = app.test_client()
    joao.post('/login', data={'organizacao': 'cliente-a', 'login': 'joao', 'senha': 'senha-segura'})
    assert b'DEST-1/2026' in joao.get('/pendencias-recebimento').data
    assert joao.post(f'/protocolo/{protocolo_id}/receber').status_code == 302
    with app.app_context():
        assert Movimentacao.query.filter_by(protocolo_id=protocolo_id).one().recebido_em is not None

        org = Organizacao.query.filter_by(slug='cliente-a').one()
        geral = Protocolo(tenant_id=org.id, numero='DEST-2/2026', nome='Destino geral',
                          data_solicitacao=date.today())
        db.session.add(geral)
        db.session.commit()
        geral_id = geral.id
    admin.post(f'/protocolo/{geral_id}/tramitar', data={'setor_destino_id': setor_id})
    assert b'DEST-2/2026' in maria.get('/pendencias-recebimento').data
    assert b'DEST-2/2026' in joao.get('/pendencias-recebimento').data
    with app.app_context():
        assert Movimentacao.query.filter_by(protocolo_id=geral_id).one().destinatario_usuario_id is None


def test_perfil_consulta_nao_pode_excluir():
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo_id = protocolo.id
    client = app.test_client()
    client.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'consulta', 'senha': 'senha-segura'
    })
    assert client.post(f'/protocolo/{protocolo_id}/deletar').status_code == 302
    with app.app_context():
        assert db.session.get(Protocolo, protocolo_id) is not None


def test_documentos_preservam_versoes_autor_historico_e_isolamento():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo_id = protocolo.id
    endpoint = f'/protocolo/{protocolo_id}/anexo/novo'
    def enviar(conteudo, documento_id=''):
        return client.post(endpoint, data={'documento_id': str(documento_id),
            'anexo': (io.BytesIO(conteudo), 'documento.txt')}, follow_redirects=True)
    assert enviar(b'primeira').status_code == 200
    with app.app_context():
        original = Anexo.query.filter_by(protocolo_id=protocolo_id).one()
        original_id, chave = original.id, original.documento_chave
    assert b'v2' in enviar(b'segunda', original_id).data
    assert b'v3' in enviar(b'terceira', original_id).data
    # Mesmo nome, mas documento independente, não deve se juntar ao anterior.
    enviar(b'independente')
    with app.app_context():
        versoes = Anexo.query.filter_by(documento_chave=chave).order_by(Anexo.versao).all()
        assert [a.versao for a in versoes] == [1, 2, 3]
        assert [a.file_data for a in versoes] == [b'primeira', b'segunda', b'terceira']
        assert all(a.enviado_por.login == 'admin' for a in versoes)
        assert Anexo.query.filter_by(protocolo_id=protocolo_id).count() == 4
        assert HistoricoProtocolo.query.filter_by(protocolo_id=protocolo_id, acao='NOVA_VERSAO_DOCUMENTO').count() == 2
    assert client.get(f'/anexo/{original_id}/download').data == b'primeira'
    client.post(f'/anexo/{original_id}/deletar')
    client.post(f'/protocolo/{protocolo_id}/deletar')
    assert client.get(f'/anexo/{original_id}/download').data == b'primeira'
    other = app.test_client()
    login(other, 'cliente-b')
    with app.app_context():
        outro_id = Protocolo.query.filter_by(nome='Dado exclusivo B').one().id
    response = other.post(f'/protocolo/{outro_id}/anexo/novo', data={
        'documento_id': str(original_id), 'anexo': (io.BytesIO(b'invasao'), 'arquivo.txt')})
    assert response.status_code == 404
    assert other.get(f'/anexo/{original_id}/download').status_code == 404


def test_novo_anexo_usa_bucket_privado_e_download_valida_tenant():
    objetos = {}
    class Body:
        def __init__(self, data): self.data = data
        def read(self): return self.data
    class FakeBucket:
        def put_object(self, Bucket, Key, Body, **kwargs): objetos[(Bucket, Key)] = Body
        def get_object(self, Bucket, Key): return {'Body': Body(objetos[(Bucket, Key)])}

    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo_id, tenant_id = protocolo.id, protocolo.tenant_id
    with patch.dict(os.environ, {'AWS_S3_BUCKET_NAME': 'bucket-teste'}), \
         patch('app.bucket_client', return_value=FakeBucket()):
        resposta = client.post(f'/protocolo/{protocolo_id}/anexo/novo', data={
            'anexo': (io.BytesIO(b'conteudo no bucket'), 'bucket.txt')},
            content_type='multipart/form-data')
        assert resposta.status_code == 302
        with app.app_context():
            anexo = Anexo.query.filter_by(protocolo_id=protocolo_id, file_name='bucket.txt').one()
            anexo_id = anexo.id
            assert anexo.storage_backend == 's3' and anexo.file_data is None
            assert anexo.storage_path.startswith(f'tenants/{tenant_id}/protocolos/{protocolo_id}/')
        download = client.get(f'/anexo/{anexo_id}/download')
        assert download.status_code == 200 and download.data == b'conteudo no bucket'

    outro = app.test_client()
    login(outro, 'cliente-b')
    assert outro.get(f'/anexo/{anexo_id}/download').status_code == 404


def test_upload_confere_conteudo_real_e_nao_apenas_extensao():
    with app.app_context():
        protocolo_id = Protocolo.query.filter_by(nome='Dado exclusivo A').one().id
        quantidade_inicial = Anexo.query.filter_by(protocolo_id=protocolo_id).count()
    client = app.test_client()
    login(client, 'cliente-a')
    endpoint = f'/protocolo/{protocolo_id}/anexo/novo'

    falso = client.post(endpoint, data={
        'anexo': (io.BytesIO(b'conteudo executavel disfarcado'), 'falso.pdf')
    }, follow_redirects=True)
    assert 'conteúdo do arquivo não corresponde' in falso.get_data(as_text=True)

    valido = client.post(endpoint, data={
        'anexo': (io.BytesIO(b'%PDF-1.4\n% teste seguro'), 'valido.pdf')
    }, follow_redirects=True)
    assert 'Documento enviado com sucesso' in valido.get_data(as_text=True)
    with app.app_context():
        assert Anexo.query.filter_by(protocolo_id=protocolo_id).count() == quantidade_inicial + 1


def test_versionamento_legado_nao_agrupa_documentos_distintos():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo_id = protocolo.id
        antigos = [Anexo(tenant_id=protocolo.tenant_id, protocolo_id=protocolo_id,
            file_name=f'legado-{i}.txt', storage_path='legado', file_size=3,
            mime_type='text/plain', file_data=b'old') for i in range(2)]
        db.session.add_all(antigos)
        db.session.commit()
        primeiro_id, segundo_id = [a.id for a in antigos]
    response = client.post(f'/protocolo/{protocolo_id}/anexo/novo', data={
        'documento_id': str(primeiro_id), 'anexo': (io.BytesIO(b'new'), 'revisao.txt')})
    assert response.status_code == 302
    with app.app_context():
        primeiro, segundo = db.session.get(Anexo, primeiro_id), db.session.get(Anexo, segundo_id)
        assert primeiro.documento_chave != segundo.documento_chave
        assert segundo.documento_chave == 'anexo'
        assert Anexo.query.filter_by(documento_chave=primeiro.documento_chave).count() == 2


def test_upload_vazio_e_processo_arquivado_nao_criam_versao():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo_id, quantidade = protocolo.id, Anexo.query.count()
    endpoint = f'/protocolo/{protocolo_id}/anexo/novo'
    assert client.post(endpoint, data={'anexo': (io.BytesIO(b''), 'vazio.txt')}).status_code == 302
    with app.app_context():
        assert Anexo.query.count() == quantidade
        protocolo = db.session.get(Protocolo, protocolo_id)
        protocolo.arquivado_em = datetime.utcnow()
        db.session.commit()
    try:
        assert client.post(endpoint, data={'anexo': (io.BytesIO(b'novo'), 'novo.txt')}).status_code == 302
        with app.app_context():
            assert Anexo.query.count() == quantidade
    finally:
        with app.app_context():
            db.session.get(Protocolo, protocolo_id).arquivado_em = None
            db.session.commit()


def test_listagem_relatorio_e_excel_usam_os_mesmos_filtros():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        alvo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        alvo.tipo_requerimento = 'Licença Especial'
        alvo.data_solicitacao = date(2026, 9, 3)
        db.session.add(Protocolo(tenant_id=alvo.tenant_id, numero='0099/2026',
            nome='Registro que deve ser excluído', tipo_requerimento='Outro',
            data_solicitacao=date(2026, 8, 1)))
        db.session.commit()
    filtros = '?nome=Dado+exclusivo+A&tipo=Licen%C3%A7a&data_inicio=2026-09-01&data_fim=2026-09-04'
    for endpoint in ['/protocolos', '/relatorios']:
        html = client.get(endpoint + filtros).get_data(as_text=True)
        assert 'Dado exclusivo A' in html
        assert 'Registro que deve ser excluído' not in html
    excel = client.get('/protocolos/backup/excel' + filtros)
    rows = list(load_workbook(io.BytesIO(excel.data)).active.iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][2] == 'Dado exclusivo A'
    assert client.get('/relatorios?data_inicio=2026-09-05&data_fim=2026-09-01').status_code == 400
    assert client.get('/protocolos?data_inicio=inválida').status_code == 400


def test_dashboard_controla_prazos_e_estatisticas_pelo_setor_atual():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        juridico = Lotacao.query.filter_by(nome='Jurídico').one()
        protocolo.setor_atual_id = juridico.id
        protocolo.prazo_em = date.today() - timedelta(days=1)
        protocolo.status = 'EM ANÁLISE'
        db.session.commit()
    dados = client.get('/protocolos/dashboard-stats?evolucaoPeriodo=all').get_json()
    assert dados['pendentesAntigos'] >= 1
    assert any(item['setor'] == 'Jurídico' and item['total'] >= 1 for item in dados['setorProtocolos'])
    filtrado = client.get('/protocolos/dashboard-stats?lotacao=Jur%C3%ADdico&evolucaoPeriodo=all').get_json()
    assert filtrado['statusProtocolos']
    assert all(item['setor'] == 'Jurídico' for item in filtrado['setorProtocolos'])
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo.status = 'ARQUIVADO'
        db.session.commit()
    assert client.get('/protocolos/dashboard-stats?evolucaoPeriodo=all').get_json()['pendentesAntigos'] == 0


def test_situacao_de_prazo_filtra_listagem_relatorio_e_excel():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        alvo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        alvo.prazo_em = date.today() + timedelta(days=4)
        alvo.status = 'EM ANÁLISE'
        fora = Protocolo.query.filter_by(nome='Registro que deve ser excluído').one()
        fora.prazo_em = date.today() + timedelta(days=15)
        fora.status = 'EM ANÁLISE'
        db.session.commit()
    filtro = '?prazo=proximos_7'
    for endpoint in ['/protocolos', '/relatorios']:
        html = client.get(endpoint + filtro).get_data(as_text=True)
        assert 'Dado exclusivo A' in html
        assert 'Registro que deve ser excluído' not in html
    rows = list(load_workbook(io.BytesIO(client.get('/protocolos/backup/excel' + filtro).data))
                .active.iter_rows(values_only=True))
    assert rows[0][16] == 'Prazo'
    assert len(rows) == 2 and rows[1][2] == 'Dado exclusivo A'
    assert client.get('/protocolos?prazo=qualquer').status_code == 400


def test_administrador_edita_desativa_e_reativa_usuario_do_cliente():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        consulta = Usuario.query.filter_by(tenant_id=1, login='consulta').one()
        juridico = Lotacao.query.filter_by(tenant_id=1, nome='Jurídico').one()
        consulta_id, juridico_id = consulta.id, juridico.id
    response = client.post(f'/admin/usuarios/{consulta_id}/editar', data={
        'nome_completo': 'Consulta Atualizada', 'email': 'consulta@example.test',
        'tipo': 'protocolista', 'lotacao_id': juridico_id})
    assert response.status_code == 302
    with app.app_context():
        consulta = db.session.get(Usuario, consulta_id)
        assert (consulta.nome_completo, consulta.tipo, consulta.lotacao_id) == ('Consulta Atualizada', 'protocolista', juridico_id)
    assert client.post(f'/admin/usuarios/{consulta_id}/status').status_code == 302
    inativo = app.test_client()
    response = inativo.post('/login', data={'organizacao': 'cliente-a', 'login': 'consulta', 'senha': 'senha-segura'}, follow_redirects=True)
    assert 'Não foi possível autenticar' in response.get_data(as_text=True)
    client.post(f'/admin/usuarios/{consulta_id}/status')
    assert b'Login bem-sucedido' in inativo.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'consulta', 'senha': 'senha-segura'}, follow_redirects=True).data


def test_usuario_de_outro_cliente_e_proprio_admin_sao_protegidos():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        admin_a = Usuario.query.filter_by(tenant_id=1, login='admin').one()
        admin_b = Usuario.query.filter_by(tenant_id=2, login='admin').one()
        admin_a_id, admin_b_id = admin_a.id, admin_b.id
    assert client.post(f'/admin/usuarios/{admin_b_id}/status').status_code == 404
    client.post(f'/admin/usuarios/{admin_a_id}/status')
    with app.app_context():
        assert db.session.get(Usuario, admin_a_id).status == 'ativo'
    client.post(f'/admin/usuarios/{admin_a_id}/editar', data={
        'nome_completo': 'Ana A', 'email': 'ana@example.test', 'tipo': 'consulta', 'lotacao_id': 0})
    with app.app_context():
        assert db.session.get(Usuario, admin_a_id).tipo == 'admin'


def test_numero_enviado_pelo_navegador_e_ignorado_e_criacao_e_atomica():
    client = app.test_client()
    login(client, 'cliente-a')
    response = client.post('/protocolo/novo', data={
        'numero': '9999/1999', 'nome': 'Numeração pelo servidor',
        'data_solicitacao': date.today().isoformat()}, follow_redirects=True)
    assert response.status_code == 200
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Numeração pelo servidor').one()
        assert protocolo.numero != '9999/1999'
        assert protocolo.numero.endswith(f'/{date.today().year}')
        assert HistoricoProtocolo.query.filter_by(protocolo_id=protocolo.id, acao='CRIACAO').count() == 1


def test_edicao_registra_campos_e_arquivado_fica_somente_leitura():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Dado exclusivo A').one()
        protocolo_id = protocolo.id
    client.post(f'/protocolo/{protocolo_id}/editar', data={
        'nome': 'Nome revisado', 'data_solicitacao': date.today().isoformat(),
        'prazo_em': (date.today() + timedelta(days=3)).isoformat()})
    with app.app_context():
        evento = HistoricoProtocolo.query.filter_by(protocolo_id=protocolo_id, acao='EDICAO').order_by(HistoricoProtocolo.id.desc()).first()
        assert 'nome:' in evento.observacao and 'prazo:' in evento.observacao
        db.session.get(Protocolo, protocolo_id).arquivado_em = datetime.utcnow()
        db.session.commit()
    client.post(f'/protocolo/{protocolo_id}/editar', data={'nome': 'Alteração proibida'})
    with app.app_context():
        assert db.session.get(Protocolo, protocolo_id).nome == 'Nome revisado'
    detalhe = client.get(f'/protocolo/{protocolo_id}').get_data(as_text=True)
    assert f'/protocolo/{protocolo_id}/editar' not in detalhe
    resposta = client.post('/protocolos/atualizar', json={'protocoloId': protocolo_id, 'novoStatus': 'STATUS LIVRE'})
    assert resposta.status_code == 400


def test_arquivamento_exige_tramitacao_regularizada():
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(nome='Nome revisado').one()
        protocolo.arquivado_em = None
        juridico = Lotacao.query.filter_by(nome='Jurídico').one()
        protocolo_id, juridico_id = protocolo.id, juridico.id
        Movimentacao.query.filter_by(protocolo_id=protocolo_id).delete()
        db.session.commit()
    client.post(f'/protocolo/{protocolo_id}/tramitar', data={'setor_destino_id': juridico_id})
    client.post(f'/protocolo/{protocolo_id}/arquivar')
    with app.app_context():
        assert db.session.get(Protocolo, protocolo_id).arquivado_em is None
    client.post(f'/protocolo/{protocolo_id}/receber')
    client.post(f'/protocolo/{protocolo_id}/arquivar')
    with app.app_context():
        protocolo = db.session.get(Protocolo, protocolo_id)
        assert protocolo.arquivado_em is not None and protocolo.status == 'ARQUIVADO'


def test_dados_institucionais_sao_isolados_e_usados_no_pdf():
    client = app.test_client()
    login(client, 'cliente-a')
    response = client.post('/admin/identidade/logo', data={
        'municipio': 'Tabuleiro do Norte/CE', 'orgao': 'Gabinete da Prefeita',
        'rodape_documento': 'Praça da Matriz — Centro', 'salvar': 'Salvar logo'},
        follow_redirects=True)
    assert response.status_code == 200 and b'Dados institucionais atualizados' in response.data
    with app.app_context():
        a = Organizacao.query.filter_by(slug='cliente-a').one()
        b = Organizacao.query.filter_by(slug='cliente-b').one()
        protocolo_id = Protocolo.query.filter_by(tenant_id=a.id).first().id
        assert (a.municipio, a.orgao) == ('Tabuleiro do Norte/CE', 'Gabinete da Prefeita')
        assert b.municipio is None and b.orgao is None
    with app.test_request_context('/'):
        with app.app_context():
            protocolo = db.session.get(Protocolo, protocolo_id)
            organizacao = Organizacao.query.filter_by(slug='cliente-a').one()
            modelo = render_template('pdf_template.html', protocolo=protocolo,
                                     organizacao=organizacao, pdf_logo_url='/logo.png')
            assert 'Gabinete da Prefeita' in modelo and 'Praça da Matriz' in modelo
            assert 'Secretaria da Administração' not in modelo
            assert 'background: #2e7d32' in modelo
            assert 'width: 52px; height: 52px' in modelo
            assert '@page { size: A4; margin: 10mm 8mm; }' in modelo
            assert 'class="paragrafo-pdf"' in modelo
            assert 'class="fechamento-pdf"' in modelo
            assert modelo.index('PROTOCOLO DE REQUERIMENTO') < modelo.index('DADOS DO REQUERENTE')
            assert modelo.index('Praça da Matriz') > modelo.index('Assinatura do Requerente')
    html = client.get(f'/protocolo/{protocolo_id}').get_data(as_text=True)
    assert 'data-organization-office="Gabinete da Prefeita"' in html
    assert 'static/img/rodape.jpg' not in html


def test_pdf_gerado_e_armazenado_com_versionamento_e_historico():
    htmls_renderizados = []

    class FakeHTML:
        def __init__(self, string, base_url):
            self.string = string
            htmls_renderizados.append(string)
        def write_pdf(self):
            return b'%PDF-1.7\nconteudo de teste'

    class FakeQRCode:
        def save(self, stream, format):
            assert format == 'PNG'
            stream.write(b'\x89PNG\r\n\x1a\nqr-de-teste')

    fake_qrcode = types.SimpleNamespace(make=lambda url: FakeQRCode())
    client = app.test_client()
    login(client, 'cliente-a')
    with app.app_context():
        protocolo = Protocolo.query.filter_by(tenant_id=1, arquivado_em=None).first()
        protocolo_id = protocolo.id
    with patch.dict(sys.modules, {'weasyprint': types.SimpleNamespace(HTML=FakeHTML),
                                  'qrcode': fake_qrcode}):
        primeira = client.post(f'/protocolo/{protocolo_id}/documento/gerar')
        segunda = client.post(f'/protocolo/{protocolo_id}/documento/gerar')
    assert primeira.status_code == 200 and primeira.content_type == 'application/pdf'
    assert segunda.status_code == 200 and '_v2.pdf' in segunda.headers['Content-Disposition']
    assert all('data:image/png;base64,' in html for html in htmls_renderizados)
    assert all('/consulta/' not in html for html in htmls_renderizados)
    with app.app_context():
        docs = Anexo.query.filter_by(protocolo_id=protocolo_id, documento_chave='documento-protocolo').order_by(Anexo.versao).all()
        assert [documento.versao for documento in docs] == [1, 2]
        assert all(documento.file_data.startswith(b'%PDF') for documento in docs)
        assert HistoricoProtocolo.query.filter_by(protocolo_id=protocolo_id, acao='DOCUMENTO_GERADO').count() == 2
    with app.app_context():
        protocolo = db.session.get(Protocolo, protocolo_id)
        protocolo.arquivado_em = datetime.utcnow()
        db.session.commit()
    with patch.dict(sys.modules, {'weasyprint': types.SimpleNamespace(HTML=FakeHTML),
                                  'qrcode': fake_qrcode}):
        assert client.post(f'/protocolo/{protocolo_id}/documento/gerar').status_code == 302
    with app.app_context():
        assert Anexo.query.filter_by(protocolo_id=protocolo_id, documento_chave='documento-protocolo').count() == 2


def test_novos_perfis_separam_protocolo_de_tramitacao():
    with app.app_context():
        tenant = Organizacao.query.filter_by(slug='cliente-a').one()
        setor = Lotacao.query.filter_by(tenant_id=tenant.id).first()
        senha = bcrypt.generate_password_hash('senha-segura').decode('utf-8')
        db.session.add_all([
            Usuario(tenant_id=tenant.id, nome='Protocolista', nome_completo='Protocolista',
                    login='protocolista', senha=senha, tipo='protocolista', lotacao_id=setor.id),
            Usuario(tenant_id=tenant.id, nome='Tramitador', nome_completo='Tramitador',
                    login='tramitador', senha=senha, tipo='tramitador', lotacao_id=setor.id),
        ])
        db.session.commit()

    protocolista = app.test_client()
    protocolista.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'protocolista', 'senha': 'senha-segura'})
    assert protocolista.get('/protocolo/novo').status_code == 200
    assert protocolista.get('/configuracoes').status_code == 302
    assert protocolista.get('/pendencias-recebimento').status_code == 200

    tramitador = app.test_client()
    tramitador.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'tramitador', 'senha': 'senha-segura'})
    assert tramitador.get('/protocolo/novo').status_code == 302
    assert tramitador.get('/configuracoes').status_code == 302
    assert tramitador.get('/pendencias-recebimento').status_code == 200


def test_tramitador_visualiza_somente_processos_do_seu_fluxo():
    with app.app_context():
        tenant = Organizacao.query.filter_by(slug='cliente-a').one()
        protocolo_setor = Lotacao.query.filter_by(tenant_id=tenant.id, nome='Protocolo').one()
        juridico = Lotacao.query.filter_by(tenant_id=tenant.id, nome='Jurídico').one()
        tramitador = Usuario.query.filter_by(tenant_id=tenant.id, login='tramitador').one()
        tramitador.lotacao_id = juridico.id
        visivel = Protocolo(tenant_id=tenant.id, numero='9001/2026', nome='Visível no Jurídico',
                            data_solicitacao=date.today(), setor_atual_id=juridico.id)
        oculto = Protocolo(tenant_id=tenant.id, numero='9002/2026', nome='Oculto no Protocolo',
                           data_solicitacao=date.today(), setor_atual_id=protocolo_setor.id)
        db.session.add_all([visivel, oculto])
        db.session.commit()
        visivel_id, oculto_id = visivel.id, oculto.id

    client = app.test_client()
    client.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'tramitador', 'senha': 'senha-segura'})
    html = client.get('/protocolos').get_data(as_text=True)
    assert 'Visível no Jurídico' in html
    assert 'Oculto no Protocolo' not in html
    assert client.get(f'/protocolo/{visivel_id}').status_code == 200
    assert client.get(f'/protocolo/{oculto_id}').status_code == 404
    assert client.get(f'/api/protocolo/{oculto_id}').status_code == 404


def test_administrador_plataforma_escolhe_cliente_explicitamente():
    with app.app_context():
        admin = Usuario.query.filter_by(tenant_id=1, login='admin').one()
        admin.is_platform_admin = True
        cliente_b = Organizacao.query.filter_by(slug='cliente-b').one()
        cliente_b_id = cliente_b.id
        db.session.commit()

    client = app.test_client()
    response = client.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura'})
    assert response.headers['Location'] == '/plataforma'
    painel = client.get('/plataforma').get_data(as_text=True)
    assert 'Cliente A' in painel and 'Cliente B' in painel
    assert client.get('/protocolos').headers['Location'] == '/plataforma'
    assert client.post(f'/plataforma/cliente/{cliente_b_id}').headers['Location'] == '/home'
    protocolos = client.get('/protocolos').get_data(as_text=True)
    assert 'Dado exclusivo B' in protocolos
    assert 'Dado exclusivo A' not in protocolos

    with app.app_context():
        admin = Usuario.query.filter_by(tenant_id=1, login='admin').one()
        admin.is_platform_admin = False
        db.session.commit()


def test_somente_administrador_geral_cria_outro_administrador_geral():
    tenant_admin = app.test_client()
    login(tenant_admin, 'cliente-a')
    assert tenant_admin.post('/plataforma/administradores/novo', data={
        'nome_completo': 'Global Bloqueado', 'login': 'global-bloqueado',
        'email': 'bloqueado@example.test', 'senha': 'senha-global', 'organizacao_id': 1,
    }).status_code == 403

    with app.app_context():
        admin = Usuario.query.filter_by(tenant_id=1, login='admin').one()
        admin.is_platform_admin = True
        cliente_b_id = Organizacao.query.filter_by(slug='cliente-b').one().id
        db.session.commit()
    global_client = app.test_client()
    global_client.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura'})
    response = global_client.post('/plataforma/administradores/novo', data={
        'nome_completo': 'Novo Global', 'login': 'novo-global',
        'email': 'global@example.test', 'senha': 'senha-global',
        'organizacao_id': cliente_b_id,
    })
    assert response.status_code == 302
    with app.app_context():
        novo = Usuario.query.filter_by(tenant_id=cliente_b_id, login='novo-global').one()
        assert novo.tipo == 'admin' and novo.is_platform_admin is True
        Usuario.query.filter_by(id=novo.id).delete()
        Usuario.query.filter_by(tenant_id=1, login='admin').one().is_platform_admin = False
        db.session.commit()


def test_administrador_do_cliente_cria_mesmo_nivel_sem_conceder_acesso_global():
    client = app.test_client()
    login(client, 'cliente-a')
    response = client.post('/admin/usuarios/novo', data={
        'nome_completo': 'Administrador Cliente', 'login': 'admin-cliente',
        'email': 'admin-cliente@example.test', 'senha': 'senha-cliente',
        'tipo': 'admin', 'lotacao_id': 0, 'is_platform_admin': 'true',
    })
    assert response.status_code == 302
    with app.app_context():
        criado = Usuario.query.filter_by(tenant_id=1, login='admin-cliente').one()
        assert criado.tipo == 'admin'
        assert criado.is_platform_admin is False
        db.session.delete(criado)
        db.session.commit()


def test_tramitador_mantem_consulta_mas_nao_altera_processo_que_saiu_do_setor():
    with app.app_context():
        tenant = Organizacao.query.filter_by(slug='cliente-a').one()
        juridico = Lotacao.query.filter_by(tenant_id=tenant.id, nome='Jurídico').one()
        protocolo_setor = Lotacao.query.filter_by(tenant_id=tenant.id, nome='Protocolo').one()
        tramitador = Usuario.query.filter_by(tenant_id=tenant.id, login='tramitador').one()
        admin = Usuario.query.filter_by(tenant_id=tenant.id, login='admin').one()
        tramitador.lotacao_id = juridico.id
        protocolo = Protocolo(tenant_id=tenant.id, numero='9003/2026', nome='Processo já encaminhado',
                              data_solicitacao=date.today(), setor_atual_id=protocolo_setor.id,
                              status='EM ANÁLISE')
        db.session.add(protocolo)
        db.session.flush()
        db.session.add(Movimentacao(
            tenant_id=tenant.id, protocolo_id=protocolo.id,
            setor_origem_id=juridico.id, setor_destino_id=protocolo_setor.id,
            enviado_por_id=tramitador.id, recebido_por_id=admin.id,
            recebido_em=datetime.utcnow()))
        db.session.commit()
        protocolo_id, destino_id = protocolo.id, juridico.id

    client = app.test_client()
    client.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'tramitador', 'senha': 'senha-segura'})
    detalhe = client.get(f'/protocolo/{protocolo_id}')
    assert detalhe.status_code == 200
    assert 'somente para consulta' in detalhe.get_data(as_text=True)
    assert client.post('/protocolos/atualizar', json={
        'protocoloId': protocolo_id, 'novoStatus': 'FINALIZADO'}).status_code == 403
    assert client.post(f'/protocolo/{protocolo_id}/tramitar', data={
        'setor_destino_id': destino_id}).status_code == 403


def test_usuario_abre_e_acompanha_apenas_os_proprios_chamados():
    consulta = app.test_client()
    consulta.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'consulta', 'senha': 'senha-segura'})
    response = consulta.post('/suporte/novo', data={
        'assunto': 'Erro ao consultar processo',
        'descricao': 'A tela informa que o processo não foi localizado.',
        'categoria': 'PROTOCOLOS', 'prioridade': 'NORMAL'}, follow_redirects=True)
    assert response.status_code == 200
    assert 'Chamado SUP-' in response.get_data(as_text=True)
    with app.app_context():
        chamado = ChamadoSuporte.query.filter_by(assunto='Erro ao consultar processo').one()
        chamado_id = chamado.id
        assert chamado.tenant_id == 1 and chamado.aberto_por.login == 'consulta'

    outro_cliente = app.test_client()
    login(outro_cliente, 'cliente-b')
    assert outro_cliente.get('/suporte').status_code == 200
    assert 'Erro ao consultar processo' not in outro_cliente.get('/suporte').get_data(as_text=True)
    assert outro_cliente.get(f'/suporte/{chamado_id}').status_code == 404

    resposta = consulta.post(f'/suporte/{chamado_id}/mensagem',
                             data={'mensagem': 'O problema continua ocorrendo.'},
                             follow_redirects=True)
    assert 'O problema continua ocorrendo.' in resposta.get_data(as_text=True)
    assert consulta.post(f'/suporte/{chamado_id}/assumir').status_code == 403
    assert consulta.post(f'/suporte/{chamado_id}/status', data={'status': 'FECHADO'}).status_code == 403


def test_administrador_geral_visualiza_assume_e_responde_chamados_de_todos_clientes():
    with app.app_context():
        admin = Usuario.query.filter_by(tenant_id=1, login='admin').one()
        admin.is_platform_admin = True
        chamado = ChamadoSuporte.query.filter_by(assunto='Erro ao consultar processo').one()
        chamado_id = chamado.id
        db.session.commit()

    global_client = app.test_client()
    global_client.post('/login', data={
        'organizacao': 'cliente-a', 'login': 'admin', 'senha': 'senha-segura'})
    painel = global_client.get('/suporte').get_data(as_text=True)
    assert 'Todos os chamados' in painel
    assert 'Erro ao consultar processo' in painel
    assert global_client.post(f'/suporte/{chamado_id}/assumir').status_code == 302
    assert global_client.post(f'/suporte/{chamado_id}/mensagem', data={
        'mensagem': 'A equipe técnica iniciou a análise.'}).status_code == 302
    assert global_client.post(f'/suporte/{chamado_id}/status', data={
        'status': 'AGUARDANDO USUÁRIO'}).status_code == 302

    with app.app_context():
        chamado = db.session.get(ChamadoSuporte, chamado_id)
        assert chamado.atribuido_a.login == 'admin'
        assert chamado.status == 'AGUARDANDO USUÁRIO'
        assert MensagemSuporte.query.filter_by(chamado_id=chamado_id).count() == 2
        Usuario.query.filter_by(tenant_id=1, login='admin').one().is_platform_admin = False
        db.session.delete(chamado)
        db.session.commit()


def test_emissao_autenticada_congela_pdf_dados_e_anexos_e_detecta_alteracao():
    htmls = []
    qr_urls = []
    class FakeHTML:
        def __init__(self, string, base_url):
            self.string = string
            htmls.append(string)
        def write_pdf(self):
            return b'%PDF-1.7\noriginal imutavel autenticado'

    class FakeQRCode:
        def save(self, stream, format):
            stream.write(b'\x89PNG\r\n\x1a\nqr')

    def fake_make(url):
        qr_urls.append(url)
        return FakeQRCode()
    fake_qrcode = types.SimpleNamespace(make=fake_make)
    with app.app_context():
        org = Organizacao.query.filter_by(slug='cliente-a').one()
        org.emissao_eletronica_protocolista_enabled = True
        org.nivel_garantia_assinatura = 'interno'
        protocolo = Protocolo(tenant_id=org.id, numero='ASS-1/2026', nome='Teste autenticado',
                              matricula='ASS-1', data_solicitacao=date.today(),
                              modalidade_abertura='presencial_protocolista')
        db.session.add(protocolo)
        db.session.commit()
        protocolo_id = protocolo.id

    client = app.test_client()
    login(client, 'cliente-a')
    modules = {'weasyprint': types.SimpleNamespace(HTML=FakeHTML), 'qrcode': fake_qrcode}
    with patch.dict(sys.modules, modules):
        response = client.post(f'/protocolo/{protocolo_id}/autenticar-emissao',
                               data={'senha': 'senha-segura'})
    assert response.status_code == 200
    assert response.data == b'%PDF-1.7\noriginal imutavel autenticado'
    assert len(qr_urls) == 2
    assert '/consulta/' in qr_urls[0] and '/validar-emissao/' in qr_urls[1]
    assert 'Acompanhe o andamento' in htmls[0] and 'Valide a autenticidade' in htmls[0]

    with app.app_context():
        emissao = EmissaoEletronica.query.filter_by(protocolo_id=protocolo_id).one()
        token = emissao.token_publico
        assert emissao.pdf_sha256 == hashlib.sha256(response.data).hexdigest()
        assert emissao.pdf_anexo.file_hash == emissao.pdf_sha256
        assert emissao.nome_emitente == 'Ana A'

    # A reimpressão recupera exatamente o original, sem regenerar nem substituir o hash.
    reprint = client.get(f'/protocolo/{protocolo_id}/pdf')
    assert reprint.data == response.data
    listagem = client.get('/protocolos').get_data(as_text=True)
    assert 'Documento autenticado' in listagem
    assert f'/protocolo/{protocolo_id}/pdf' in listagem
    assert client.post(f'/protocolo/{protocolo_id}/editar', data={'nome': 'Alterado'}).status_code == 302
    assert client.post(f'/protocolo/{protocolo_id}/anexo/novo', data={}).status_code == 302

    tela_retificacao = client.get(f'/protocolo/{protocolo_id}/retificar').get_data(as_text=True)
    assert 'Retificar Protocolo ASS-1/2026' in tela_retificacao
    criada = client.post(f'/protocolo/{protocolo_id}/retificar', data={
        'nome': 'Teste autenticado retificado', 'matricula': 'ASS-1',
        'data_solicitacao': date.today().isoformat(), 'observacoes': 'Correção formal.'})
    assert criada.status_code == 302
    with app.app_context():
        retificacao = db.session.get(Protocolo, protocolo_id)
        assert retificacao.numero == 'ASS-1/2026'
        assert retificacao.retificacao_pendente is True
        assert len(retificacao.emissoes_eletronicas) == 1
    qr_urls.clear()
    with patch.dict(sys.modules, modules):
        nova_emissao = client.post(f'/protocolo/{protocolo_id}/autenticar-emissao',
                                   data={'senha': 'senha-segura'})
    assert nova_emissao.status_code == 200 and len(qr_urls) == 2
    with app.app_context():
        retificacao = db.session.get(Protocolo, protocolo_id)
        assert retificacao.numero == 'ASS-1/2026'
        assert [e.versao for e in retificacao.emissoes_eletronicas] == [1, 2]
        assert [e.status for e in retificacao.emissoes_eletronicas] == ['RETIFICADA', 'VALIDA']
        assert retificacao.emissao_eletronica.versao == 2
        token_vigente = retificacao.emissao_eletronica.token_publico
    detalhe_versoes = client.get(f'/protocolo/{protocolo_id}').get_data(as_text=True)
    assert 'Versão 2' in detalhe_versoes and 'Versão 1' in detalhe_versoes

    # Cancelamento exige credencial e motivo, preserva o PDF e torna pública a situação.
    motivo_curto = client.post(f'/protocolo/{protocolo_id}/cancelar-emissao', data={
        'senha': 'senha-segura', 'motivo': 'erro'}, follow_redirects=True)
    assert 'pelo menos 10 caracteres' in motivo_curto.get_data(as_text=True)
    cancelada = client.post(f'/protocolo/{protocolo_id}/cancelar-emissao', data={
        'senha': 'senha-segura', 'motivo': 'Pedido cancelado formalmente pelo setor responsável.'},
        follow_redirects=True)
    assert 'foi cancelada sem apagar' in cancelada.get_data(as_text=True)
    with app.app_context():
        retificacao = db.session.get(Protocolo, protocolo_id)
        assert retificacao.emissao_eletronica.status == 'CANCELADA'
        assert retificacao.emissao_eletronica.cancelado_por_id is not None
        assert retificacao.emissao_eletronica.cancelado_em is not None
        assert len(retificacao.emissoes_eletronicas) == 2
        assert HistoricoProtocolo.query.filter_by(
            protocolo_id=protocolo_id, acao='EMISSAO_CANCELADA').count() == 1
    assert client.get(f'/protocolo/{protocolo_id}/pdf').data == nova_emissao.data
    consulta_cancelada = client.get(f'/validar-emissao/{token_vigente}').get_data(as_text=True)
    assert 'esta emissão foi cancelada' in consulta_cancelada
    assert 'Pedido cancelado formalmente' in consulta_cancelada
    assert client.get(f'/protocolo/{protocolo_id}/retificar').status_code == 409

    valido = client.post(f'/validar-emissao/{token}', data={
        'arquivo': (io.BytesIO(response.data), 'original.pdf')},
        content_type='multipart/form-data').get_data(as_text=True)
    adulterado = client.post(f'/validar-emissao/{token}', data={
        'arquivo': (io.BytesIO(response.data + b'alteracao'), 'alterado.pdf')},
        content_type='multipart/form-data').get_data(as_text=True)
    assert 'Integridade confirmada' in valido
    assert 'Integridade não confirmada' in adulterado

    with app.app_context():
        org = Organizacao.query.filter_by(slug='cliente-a').one()
        org.emissao_eletronica_protocolista_enabled = False
        db.session.commit()


def test_portal_servidor_isola_login_e_abertura_em_nome_proprio():
    with app.app_context():
        org = Organizacao.query.filter_by(slug='cliente-a').one()
        org.portal_servidor_remoto_enabled = True
        servidor = Servidor(tenant_id=org.id, matricula='PORTAL-1', nome='Servidor Portal',
                            cargo='Analista', lotacao='Protocolo', unidade_de_exercicio='Sede')
        tipo = TipoRequerimento(tenant_id=org.id, nome='Requerimento remoto', ativo=True)
        db.session.add_all([servidor, tipo])
        db.session.flush()
        senha = bcrypt.generate_password_hash('senha-portal').decode('utf-8')
        usuario = Usuario(tenant_id=org.id, nome='Servidor', nome_completo='Servidor Portal',
                          login='servidor.portal', senha=senha, tipo='requerente',
                          servidor_id=servidor.id, status='ativo')
        db.session.add(usuario)
        db.session.commit()
        servidor_id = servidor.id

    client = app.test_client()
    # A credencial do requerente não pode entrar pela superfície administrativa.
    backoffice = client.post('/entrar/cliente-a', data={
        'login': 'servidor.portal', 'senha': 'senha-portal'})
    assert backoffice.status_code == 200
    portal = client.post('/portal/cliente-a/entrar', data={
        'login': 'servidor.portal', 'senha': 'senha-portal'})
    assert portal.status_code == 302 and portal.headers['Location'] == '/portal'
    home = client.get('/portal').get_data(as_text=True)
    assert 'Olá, Servidor Portal' in home and 'Novo requerimento' in home

    portal_html = []
    class PortalHTML:
        def __init__(self, string, base_url):
            self.string = string
            portal_html.append(string)
        def write_pdf(self):
            return b'%PDF-1.7\nenvio remoto autenticado'
    class PortalQR:
        def save(self, stream, format):
            stream.write(b'\x89PNG\r\n\x1a\nqr')
    with patch.dict(sys.modules, {
            'weasyprint': types.SimpleNamespace(HTML=PortalHTML),
            'qrcode': types.SimpleNamespace(make=lambda _url: PortalQR())}):
        criado = client.post('/portal/novo', data={
            'tipo_requerimento': 'Requerimento remoto', 'requer_ao': 'Setor responsável',
            'observacoes': 'Solicito análise do pedido enviado de casa.', 'declaracao': 'on'})
    assert criado.status_code == 302
    assert 'REQUERIMENTO ENVIADO ELETRONICAMENTE' in portal_html[0]
    assert 'Requerente: Servidor Portal' in portal_html[0]
    assert 'conta individual vinculada ao cadastro funcional' not in portal_html[0]
    assert portal_html[0].count('Autenticidade verificável pelo QR Code.') == 1
    with app.app_context():
        protocolo = Protocolo.query.filter_by(requerente_servidor_id=servidor_id).one()
        assert protocolo.nome == 'Servidor Portal'
        assert protocolo.matricula == 'PORTAL-1'
        assert protocolo.modalidade_abertura == 'remota_requerente'
        assert protocolo.emissao_eletronica.metodo == 'conta_individual'
        assert protocolo.emissao_eletronica.pdf_sha256 == hashlib.sha256(
            b'%PDF-1.7\nenvio remoto autenticado').hexdigest()
        protocolo_id = protocolo.id
    detalhe = client.get(f'/portal/protocolo/{protocolo_id}').get_data(as_text=True)
    assert 'Solicito análise do pedido enviado de casa.' in detalhe
    # Rotas do backoffice não ficam disponíveis na sessão do portal.
    assert client.get('/configuracoes').status_code == 302
    assert client.get('/portal/protocolo/1').status_code == 404

    client.post('/portal/sair')
    with app.app_context():
        Protocolo.query.filter_by(requerente_servidor_id=servidor_id).delete()
        Usuario.query.filter_by(login='servidor.portal').delete()
        Servidor.query.filter_by(id=servidor_id).delete()
        TipoRequerimento.query.filter_by(nome='Requerimento remoto').delete()
        Organizacao.query.filter_by(slug='cliente-a').one().portal_servidor_remoto_enabled = False
        db.session.commit()

