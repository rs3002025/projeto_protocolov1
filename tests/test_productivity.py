"""Funcionalidades de produtividade com banco e identidades fictícios."""
import io,json
from datetime import date,datetime,timedelta,timezone
from unittest import TestCase
from unittest.mock import patch
import test_multitenancy as _isolated_database_setup
import test_protocol_workflow as fixture
from flask import template_rendered
from sqlalchemy import event
from sqlalchemy.dialects import postgresql
from openpyxl import load_workbook
from app import app,db
from models import (PesquisaSalva,ModeloDespacho,PreferenciaAlerta,Protocolo,Movimentacao,
                    ProvidenciaProtocolo,DespachoProtocolo,Usuario)
from management_reports import management_data


class ProductivityTests(TestCase):
    setUp=fixture.WorkflowTests.setUp
    tearDown=fixture.WorkflowTests.tearDown
    login=fixture.WorkflowTests.login
    post=fixture.WorkflowTests.post
    send=fixture.WorkflowTests.send
    receive=fixture.WorkflowTests.receive
    participation=fixture.WorkflowTests.participation

    def context(self,path,template_name):
        contexts=[]
        def capture(sender,template,context,**extra):
            if template.name==template_name: contexts.append(context)
        with template_rendered.connected_to(capture,app): response=self.client.get(path)
        self.assertEqual(response.status_code,200,response.get_data(as_text=True)[:500])
        return contexts[0]

    def test_searches_are_personal_scoped_validated_and_reversible(self):
        self.login('admin')
        response=self.client.post('/relatorios/pesquisas',data={'nome_pesquisa':'Minha pesquisa','setor_atual':'SEFIN','prazo':'vencido','next':'https://example.org'})
        self.assertEqual(response.status_code,302)
        with app.app_context():
            item=PesquisaSalva.query.one(); identifier=item.id
            self.assertNotIn('next',json.loads(item.filtros))
        context=self.context('/relatorios','relatorios.html')
        self.assertEqual(len(context['pesquisas_salvas']),1)
        self.assertIn('setor_atual=SEFIN',context['pesquisas_salvas'][0]['url'])
        self.login('protocolista')
        self.assertEqual(self.context('/relatorios','relatorios.html')['pesquisas_salvas'],[])
        self.assertEqual(self.client.post(f'/relatorios/pesquisas/{identifier}/remover').status_code,404)
        self.login('foreign'); self.assertEqual(self.client.post(f'/relatorios/pesquisas/{identifier}/remover',headers={'Content-Type':'application/json'}).status_code,403)
        self.login('admin')
        self.assertEqual(self.client.post('/relatorios/pesquisas',data={'nome_pesquisa':'Inválida','data_inicio':'bad'}).status_code,400)
        self.assertEqual(self.client.post(f'/relatorios/pesquisas/{identifier}/remover').status_code,302)
        self.assertEqual(self.context('/relatorios','relatorios.html')['pesquisas_salvas'],[])
        self.client.post('/relatorios/pesquisas',data={'nome_pesquisa':'Minha pesquisa','status':'EM ANÁLISE'})
        with app.app_context(): self.assertEqual(PesquisaSalva.query.count(),1); self.assertTrue(PesquisaSalva.query.one().ativo)

    def test_templates_keep_formatting_and_do_not_register_dispatch_automatically(self):
        self.login('admin')
        delta=json.dumps({'ops':[{'insert':'Modelo revisável','attributes':{'bold':True}},{'insert':'\n'}]})
        self.assertEqual(self.client.post('/modelos-despacho/salvar',data={'nome_modelo':'Parecer','assunto':'Análise','texto':'x','conteudo_delta':delta}).status_code,302)
        with app.app_context(): identifier=ModeloDespacho.query.one().id
        self.login('protocolista')
        context=self.context(f'/protocolo/{self.pid}/despachos/novo?modelo_id={identifier}','despacho_editor.html')
        self.assertEqual(context['assunto'],'Análise'); self.assertIn('bold',context['conteudo_delta'])
        with app.app_context(): self.assertEqual(DespachoProtocolo.query.count(),0)
        self.assertEqual(self.client.post('/modelos-despacho/salvar',headers={'Content-Type':'application/json'}).status_code,403)
        self.login('foreign'); self.assertEqual(self.client.get(f'/protocolo/{self.foreign_pid}/despachos/novo?modelo_id={identifier}').status_code,404)
        self.login('admin')
        self.assertEqual(self.client.post(f'/modelos-despacho/{identifier}/situacao',data={'ativo':'0'}).status_code,302)
        self.assertEqual(self.client.get(f'/protocolo/{self.pid}/despachos/novo?modelo_id={identifier}').status_code,404)
        self.assertEqual(self.client.post(f'/modelos-despacho/{identifier}/situacao',data={'ativo':'1'}).status_code,302)
        self.assertEqual(self.client.get(f'/protocolo/{self.pid}/despachos/novo?modelo_id={identifier}').status_code,200)
        self.assertEqual(self.client.get('/modelos-despacho/editar').status_code,200)

    def test_alert_settings_are_personal_and_deadlines_exclude_closed(self):
        self.login('admin'); today=date(2026,10,24)
        with app.app_context():
            db.session.get(Protocolo,self.pid).prazo_em=today+timedelta(days=2)
            db.session.add(Protocolo(tenant_id=self.tenant,numero='FECHADO',nome='Fechado',status='CONCLUÍDO',prazo_em=today-timedelta(days=2))); db.session.commit()
        with patch('dashboard_workbench.today_brasilia',return_value=today):
            wb=self.context('/home?fila=alertas','home.html')['workbench']
            self.assertEqual(wb['counts']['alertas'],1)
            self.client.post('/painel/alertas',data={'ativo':'1','antecedencia_dias':'1'})
            self.assertEqual(self.context('/home?fila=alertas','home.html')['workbench']['counts']['alertas'],0)
            self.login('protocolista')
            self.assertEqual(self.context('/home?fila=alertas','home.html')['workbench']['counts']['alertas'],1)
        self.assertEqual(self.client.post('/painel/alertas',data={'antecedencia_dias':'100'}).status_code,400)
        with app.app_context(): self.assertEqual(PreferenciaAlerta.query.count(),1)

    def test_management_separates_receipt_wait_service_and_pending_transfers(self):
        self.login('admin'); self.send('alice'); self.receive('alice')
        old_part,old_move=self.participation('alice')
        self.send('carol',providencia_id=old_part)
        _,new_move=self.participation('carol')
        self.login('admin')
        with app.app_context():
            first=db.session.get(Movimentacao,old_move); second=db.session.get(Movimentacao,new_move)
            first.enviado_em=datetime(2026,10,1,12); first.recebido_em=datetime(2026,10,1,14)
            second.enviado_em=datetime(2026,10,2,14)
            p=db.session.get(Protocolo,self.pid); p.data_solicitacao=date(2026,10,1)
            db.session.commit()
        report=self.context('/relatorios/gestao?data_inicio=2026-10-01&data_fim=2026-10-03','relatorio_gestao.html')['report']
        first=next(r for r in report['sectors'] if r['nome']=='SEFIN')
        self.assertEqual(first['espera_media'],2)
        self.assertEqual(first['atendimento_medio'],24)
        self.assertEqual(first['pendentes'],0)
        self.assertEqual(sum(r['pendentes'] for r in report['sectors']),1)
        self.assertEqual(report['entradas'],1)
        # A saída seguinte fica fora do recorte, mas encerra o atendimento.
        earlier=self.context('/relatorios/gestao?data_inicio=2026-10-01&data_fim=2026-10-01','relatorio_gestao.html')['report']
        self.assertEqual(earlier['sectors'][0]['atendimento_medio'],24)
        self.assertEqual(self.client.get('/relatorios/gestao?data_inicio=bad').status_code,400)
        self.assertEqual(self.client.get('/relatorios/gestao?data_inicio=1992-01-01&data_fim=2026-01-01').status_code,400)

    def test_closing_uses_brasilia_date_and_unknown_dates_stay_out_of_average(self):
        self.login('admin')
        with app.app_context():
            p=db.session.get(Protocolo,self.pid); p.data_solicitacao=date(2026,10,1)
            p.status='CONCLUÍDO'; p.finalizado_em=datetime(2026,10,3,2) # 02/10 às 23h Brasília.
            db.session.add(Protocolo(tenant_id=self.tenant,numero='SEM-DATA',nome='Sem data',status='CONCLUÍDO'))
            db.session.commit()
        report=self.context('/relatorios/gestao?data_inicio=2026-10-02&data_fim=2026-10-02','relatorio_gestao.html')['report']
        self.assertEqual(report['conclusoes'],1); self.assertEqual(report['media_dias'],1)
        self.assertEqual(report['sem_data_conclusao'],1)
        data=self.client.get('/relatorios/gestao/excel?data_inicio=2026-10-02&data_fim=2026-10-02')
        self.assertEqual(data.status_code,200)
        workbook=load_workbook(io.BytesIO(data.data)); self.assertEqual(len(workbook.sheetnames),3)
        self.assertEqual(workbook['Entradas e conclusões'].freeze_panes,'A3')
        self.assertTrue(workbook['Entradas e conclusões'].auto_filter.ref.startswith('A2:'))
        self.login('alice'); self.assertEqual(self.client.get('/relatorios/gestao',headers={'Content-Type':'application/json'}).status_code,403)
        self.assertEqual(self.client.get('/relatorios/gestao/excel',headers={'Content-Type':'application/json'}).status_code,403)

    def test_management_sql_compiles_for_postgresql(self):
        self.login('admin'); self.send('alice')
        statements=[]
        def capture(conn,cursor,statement,parameters,context,executemany):
            if context.compiled: statements.append(str(context.compiled.statement.compile(dialect=postgresql.dialect())))
        with app.app_context(): engine=db.engine; event.listen(engine,'before_cursor_execute',capture)
        try: self.context('/relatorios/gestao','relatorio_gestao.html')
        finally: event.remove(engine,'before_cursor_execute',capture)
        self.assertTrue(any('lead(' in sql for sql in statements))

    def test_foreign_administrator_cannot_edit_or_apply_another_clients_template(self):
        self.login('admin')
        self.client.post('/modelos-despacho/salvar',data={'nome_modelo':'Privado do cliente A','assunto':'Teste','texto':'Modelo fictício'})
        with app.app_context():
            identifier=ModeloDespacho.query.one().id
            db.session.get(Usuario,self.ids['foreign']).tipo='admin'; db.session.commit()
        self.login('foreign')
        self.assertEqual(self.client.get(f'/modelos-despacho/{identifier}/editar').status_code,404)
        self.assertEqual(self.client.get(f'/protocolo/{self.foreign_pid}/despachos/novo?modelo_id={identifier}').status_code,404)
        self.assertNotIn('Privado do cliente A',self.client.get('/modelos-despacho').get_data(as_text=True))

    def test_superseded_legacy_pending_is_not_counted_as_current(self):
        self.login('admin')
        with app.app_context():
            db.session.add(Movimentacao(tenant_id=self.tenant,protocolo_id=self.pid,setor_destino_id=self.sector,
                enviado_por_id=self.ids['admin'],enviado_em=datetime(2026,10,1,12)))
            db.session.flush()
            participation=ProvidenciaProtocolo(tenant_id=self.tenant,protocolo_id=self.pid)
            db.session.add(participation); db.session.flush()
            db.session.add(Movimentacao(tenant_id=self.tenant,protocolo_id=self.pid,setor_destino_id=self.sector,
                providencia_id=participation.id,enviado_por_id=self.ids['admin'],enviado_em=datetime(2026,10,3,12)))
            db.session.commit()
        report=self.context('/relatorios/gestao?data_inicio=2026-10-01&data_fim=2026-10-05','relatorio_gestao.html')['report']
        self.assertEqual(sum(row['pendentes'] for row in report['sectors']),1)
        self.assertEqual(sum(row['redirecionados'] for row in report['sectors']),1)
