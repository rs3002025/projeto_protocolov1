"""Reproduções isoladas de problemas; não executa contra Railway."""
import io
import zipfile
from datetime import datetime
import test_multitenancy as base
from app import app, db, login_fingerprint, verify_audit_chain, verified_upload_mime
from models import Protocolo, HistoricoProtocolo, EmissaoEletronica, Anexo
from openpyxl import load_workbook

def setup_module():
    from pathlib import Path
    Path(base._test_directory.name).mkdir(exist_ok=True)
    base.setup_module()

def teardown_module():
    base.teardown_module()

def test_timestamp_tamper_is_detected():
    with app.app_context():
        h = HistoricoProtocolo(tenant_id=1, protocolo_id=1, acao='AUDIT_TEST', status='TESTE', responsavel='admin')
        db.session.add(h)
        db.session.commit()
        assert verify_audit_chain(1)[0]
        h.data_movimentacao = datetime(2000, 1, 1)
        db.session.commit()
        assert not verify_audit_chain(1)[0]

def test_forwarded_prefix_changes_limiter_identity():
    with app.test_request_context('/', headers={'X-Forwarded-For':'attacker-one, 203.0.113.1'}):
        a = login_fingerprint('cliente-a', 'admin')
    with app.test_request_context('/', headers={'X-Forwarded-For':'attacker-two, 203.0.113.1'}):
        b = login_fingerprint('cliente-a', 'admin')
    assert a == b

def test_excel_exports_user_text_as_formula():
    with app.app_context():
        db.session.get(Protocolo, 1).observacoes = '=1+1'
        db.session.commit()
    c = app.test_client()
    base.login(c, 'cliente-a')
    r = c.get('/protocolos/backup/excel')
    assert r.status_code == 200
    assert load_workbook(io.BytesIO(r.data)).active['R2'].data_type == 's'

def test_archived_authenticated_protocol_can_be_rectified():
    with app.app_context():
        p = db.session.get(Protocolo, 1)
        p.arquivado_em = datetime.now()
        a = Anexo(tenant_id=1, protocolo_id=1, file_name='test.pdf', storage_path='test', file_size=4, mime_type='application/pdf', file_data=b'test')
        db.session.add(a)
        db.session.flush()
        db.session.add(EmissaoEletronica(tenant_id=1, protocolo_id=1, emitido_por_id=1, pdf_anexo_id=a.id, token_publico='audit-token', codigo_publico='audit-code', pdf_sha256='0'*64, declaracao='Teste', nome_emitente='Admin', login_emitente='admin', ip_hash='0'*64, user_agent_hash='0'*64))
        db.session.commit()
    c = app.test_client()
    base.login(c, 'cliente-a')
    r = c.post('/protocolo/1/retificar', data={'nome':'ALTERADO ARQUIVADO'})
    assert r.status_code == 409
    with app.app_context():
        p = db.session.get(Protocolo, 1)
        assert p.nome != 'ALTERADO ARQUIVADO' and not p.retificacao_pendente

def test_renamed_zip_is_rejected_and_real_workbook_accepted():
    from openpyxl import Workbook
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w') as archive:
        archive.writestr('arquivo.txt', 'Teste')
    assert verified_upload_mime('arquivo.docx', stream.getvalue()) is None
    assert verified_upload_mime('arquivo.xlsx', stream.getvalue()) is None
    workbook = io.BytesIO()
    Workbook().save(workbook)
    assert verified_upload_mime('arquivo.xlsx', workbook.getvalue()) is not None

def test_password_change_revokes_other_sessions():
    from models import Usuario
    first, second = app.test_client(), app.test_client()
    base.login(first, 'cliente-a')
    base.login(second, 'cliente-a')
    assert second.get('/protocolos').status_code == 200
    response = first.post('/minha-conta/trocar-senha', data={
        'senha_atual': 'senha-segura', 'senha_nova': 'NovaSenhaSegura123',
        'confirmar_senha': 'NovaSenhaSegura123'})
    assert response.status_code == 302
    assert second.get('/protocolos').status_code == 302
    assert first.get('/protocolos').status_code == 200
    with app.app_context():
        Usuario.query.filter_by(tenant_id=1, login='admin').one().senha = base.bcrypt.generate_password_hash('senha-segura').decode()
        db.session.commit()

