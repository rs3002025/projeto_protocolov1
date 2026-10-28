"""Exercita ferramentas reais em PostgreSQL descartável local. Não aceita banco remoto."""
import os
import sys
import io
import json
import base64
import zipfile
from pathlib import Path
from datetime import date,datetime
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urlsplit
from unittest.mock import patch

url=os.environ['QA_DATABASE_URL']
parsed=urlsplit(url)
if parsed.hostname not in ('localhost','127.0.0.1') or parsed.path!='/muniprot_qa' or parsed.username!='qa':
    raise RuntimeError('Somente PostgreSQL local descartável muniprot_qa, usuário qa.')
os.environ.update(DATABASE_URL=url,SECRET_KEY='disposable-functional-qa',TENANT_SUBDOMAINS_ENABLED='false')
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import app,db,bcrypt,verify_audit_chain
import app as application
from models import (Organizacao,Usuario,Lotacao,TipoRequerimento,Servidor,Protocolo,
                    HistoricoProtocolo,Anexo,Movimentacao,ProvidenciaProtocolo,
                    ModeloDespacho,DespachoProtocolo,PesquisaSalva,ChamadoSuporte)
from openpyxl import Workbook,load_workbook
from pypdf import PdfReader
from PIL import Image
from bootstrap_db import align_server_sequence
from sqlalchemy import text

app.config.update(TESTING=True,WTF_CSRF_ENABLED=False)
passed=[]
def check(condition,name):
    if not condition: raise AssertionError(name)
    passed.append(name)

def request(client,method,path,expected=200,**kwargs):
    response=getattr(client,method)(path,**kwargs)
    expected=(expected,) if isinstance(expected,int) else expected
    check(response.status_code in expected,f'{method.upper()} {path}: {response.status_code}, esperado {expected}')
    return response

def row(model,**filters):
    with app.app_context():
        item=model.query.filter_by(**filters).one()
        return {c.name:getattr(item,c.name) for c in model.__table__.columns}

def client_for(name):
    client=app.test_client()
    if name=='global':
        request(client,'post','/login',302,data={'login':'global','senha':password})
        request(client,'post',f'/plataforma/cliente/{tenant}',302)
    else:
        request(client,'post','/entrar/qa-a',302,data={'login':name,'senha':password})
    return client

password='Ficticio-Postgres-2026'
png=io.BytesIO();Image.new('RGB',(24,24),'green').save(png,format='PNG')
logo='data:image/png;base64,'+base64.b64encode(png.getvalue()).decode()
real_public_url=application.public_url
def fixture_asset_url(endpoint,**kwargs):
    return logo if endpoint=='organization_logo' else real_public_url(endpoint,**kwargs)

with app.app_context():
    assert db.engine.dialect.name=='postgresql'
    db.create_all()
    a=Organizacao(nome='QA A',slug='qa-a',portal_servidor_remoto_enabled=True)
    b=Organizacao(nome='QA B',slug='qa-b')
    db.session.add_all([a,b]);db.session.flush()
    tenant,other=a.id,b.id
    sectors=[Lotacao(tenant_id=tenant,nome=n,ativo=True) for n in ('QA SEAD','QA SEFIN','QA CULTURA')]
    db.session.add_all(sectors);db.session.flush()
    sector_ids=[s.id for s in sectors]
    hashed=bcrypt.generate_password_hash(password).decode()
    users={}
    for name,role,sector in [('admin','admin',sector_ids[0]),('protocolo','protocolista',sector_ids[0]),
                             ('alice','tramitador',sector_ids[0]),('bob','tramitador',sector_ids[1]),
                             ('carol','tramitador',sector_ids[2]),('consulta','consulta',sector_ids[0])]:
        item=Usuario(tenant_id=tenant,nome='QA '+name,nome_completo='Pessoa fictícia '+name,
                     login=name,senha=hashed,tipo=role,lotacao_id=sector,status='ativo')
        db.session.add(item);db.session.flush();users[name]={'id':item.id,'session':item.get_id()}
    global_user=Usuario(nome='QA Global',login='global',senha=hashed,tipo='admin',is_platform_admin=True)
    foreign_user=Usuario(tenant_id=other,nome='QA Externo',login='foreign',senha=hashed,tipo='admin')
    server=Servidor(tenant_id=tenant,matricula='QA-PORTAL',nome='Requerente fictício',cpf='01234567890',
                    nascimento=date(1990,1,1),nome_mae='Mãe fictícia',email='qa@example.test')
    db.session.add_all([global_user,foreign_user,server]);db.session.flush()
    portal_user=Usuario(tenant_id=tenant,nome=server.nome,login=server.matricula,senha=hashed,tipo='requerente',
                        servidor_id=server.id,email='qa@example.test',status='ativo')
    db.session.add_all([portal_user,TipoRequerimento(tenant_id=tenant,nome='QA Pedido',ativo=True)])
    foreign_protocol=Protocolo(tenant_id=other,numero='0001/2026',nome='Exclusivo do outro cliente',status='RECEBIDO')
    db.session.add(foreign_protocol);db.session.commit()
    foreign_id=foreign_protocol.id
    # Reproduz a carga histórica com identificadores explícitos e sequência atrasada.
    db.session.add(Servidor(id=50000,tenant_id=tenant,matricula='QA-LEGADO',nome='Legado fictício'))
    db.session.commit()
    with db.engine.begin() as connection:
        align_server_sequence(connection)
        first=connection.execute(text("SELECT pg_sequence_last_value(pg_get_serial_sequence('servidores','id')::regclass)")).scalar_one()
        check(first>=50000,'Carga histórica: sequência avança além dos IDs existentes')
        connection.execute(text("SELECT setval(pg_get_serial_sequence('servidores','id'),60000,true)"))
        align_server_sequence(connection)
        check(connection.execute(text("SELECT pg_sequence_last_value(pg_get_serial_sequence('servidores','id')::regclass)")).scalar_one()==60000,'Sequência adiantada nunca é reduzida')

try:
  with patch('app.public_url',side_effect=fixture_asset_url), patch('app.bucket_client',return_value=None):
    admin=client_for('admin');protocolist=client_for('protocolo');viewer=client_for('consulta')
    alice=client_for('alice');bob=client_for('bob');carol=client_for('carol');global_client=client_for('global')
    for path in ('/home','/protocolos','/meus_protocolos','/relatorios','/relatorios/gestao','/configuracoes',
                 '/modelos-despacho','/admin/recuperar-documento','/admin/auditoria','/admin/servidores/cadastro-portal'):
        request(admin,'get',path)
    for kind,model in [('tipo',TipoRequerimento),('lotacao',Lotacao)]:
        for name in ('Cadastro QA','Cadastro QA',' cadastro qa '):
            request(admin,'post',f'/admin/item/{kind}/novo',302,data={'nome':name})
        item=row(model,tenant_id=tenant,nome='Cadastro QA')
        request(admin,'post',f'/admin/item/{kind}/{item["id"]}/status',302)
        request(admin,'post',f'/admin/item/{kind}/novo',302,data={'nome':'Cadastro QA'})
        check(not row(model,id=item['id'])['ativo'],f'{kind}: repetição não reativa item inativo')
        with app.app_context():check(model.query.filter_by(tenant_id=tenant,nome='Cadastro QA').count()==1,f'{kind}: sem duplicatas')

    model_file=request(admin,'get','/admin/servidores/modelo')
    workbook=load_workbook(io.BytesIO(model_file.data));sheet=workbook['Servidores']
    headers=[c.value for c in sheet[1]]
    for column in ('matricula','nome','cpf','data_nascimento','nome_mae','endereco','telefone','email'):
        check(column in headers,f'Planilha-modelo contém {column}')
    workbook=Workbook();sheet=workbook.active;sheet.title='Servidores'
    sheet.append(['matricula','nome','cpf','data_nascimento','nome_mae','endereco','telefone','email'])
    sheet.append(['QA-IMPORT','Pessoa importada fictícia','12345678901','01/01/1990','Mãe fictícia','Endereço fictício','00000000000','qa@example.test'])
    data=io.BytesIO();workbook.save(data)
    for _ in range(2):
        request(admin,'post','/admin/servidores/importar',302,data={'arquivo':(io.BytesIO(data.getvalue()),'qa.xlsx')},content_type='multipart/form-data')
    check(row(Servidor,tenant_id=tenant,matricula='QA-IMPORT')['telefone']=='00000000000','Importação e reimportação preservam contato')

    def create_one(index):
        client=app.test_client()
        with client.session_transaction() as session:session['_user_id'],session['_fresh']=users['admin']['session'],True
        return request(client,'post','/protocolo/novo',302,data={'nome':f'Concorrência QA {index}',
            'data_solicitacao':'2026-10-28','tipo_requerimento':'QA Pedido'}).status_code
    with ThreadPoolExecutor(max_workers=6) as pool:list(pool.map(create_one,range(6)))
    with app.app_context():
        created=Protocolo.query.filter_by(tenant_id=tenant).all()
        check(len(created)==6 and len({p.numero for p in created})==6,'Seis cadastros simultâneos têm numeração única')
        pid=created[0].id
    path=f'/protocolo/{pid}'
    request(admin,'post','/protocolos/atualizar',json={'protocoloId':str(pid),'novoStatus':'EM ANÁLISE'})
    request(viewer,'post','/protocolos/atualizar',403,json={'protocoloId':str(pid),'novoStatus':'EM ANÁLISE'})
    request(admin,'get',f'/protocolo/{foreign_id}',404)
    request(admin,'post','/modelos-despacho/salvar',302,data={'nome_modelo':'Modelo QA','assunto':'Parecer QA','texto':'Texto fictício do modelo.'})
    model=row(ModeloDespacho,tenant_id=tenant,nome='Modelo QA')
    request(protocolist,'get',f'{path}/despachos/novo?modelo_id={model["id"]}')
    request(admin,'post',f'{path}/despachos',302,data={'assunto':'Despacho QA','texto':'Texto formal fictício.\n\nSegundo parágrafo.', 'envio_token':'postgres-dispatch-token-00001'})
    dispatch=row(DespachoProtocolo,protocolo_id=pid,numero=1)
    pdf=request(admin,'get',f'{path}/despachos/{dispatch["id"]}/pdf')
    check(len(PdfReader(io.BytesIO(pdf.data)).pages)>=1,'Despacho renderizado como PDF real')
    request(admin,'post',f'{path}/despachos',302,data={'assunto':'Despacho QA','texto':'Texto repetido','envio_token':'postgres-dispatch-token-00001'})
    with app.app_context():check(DespachoProtocolo.query.filter_by(protocolo_id=pid).count()==1,'Reenvio do despacho é idempotente')

    original=b'Anexo ficticio versao 1.';replacement=b'Anexo ficticio versao 2.'
    request(admin,'post',f'{path}/anexo/novo',302,data={'anexo':(io.BytesIO(original),'qa.txt')},content_type='multipart/form-data')
    attachment=row(Anexo,protocolo_id=pid,versao=1)
    request(admin,'post',f'{path}/anexo/novo',302,data={'documento_id':str(attachment['id']),'anexo':(io.BytesIO(replacement),'qa-v2.txt')},content_type='multipart/form-data')
    current=row(Anexo,protocolo_id=pid,versao=2)
    check(request(admin,'get',f'/anexo/{current["id"]}/download').data==replacement,'Download da substituição preserva bytes')
    request(admin,'post',f'/admin/recuperar-documento/{current["id"]}',409,data={'original':(io.BytesIO(replacement),'qa.txt')},content_type='multipart/form-data')
    with app.app_context():
        item=db.session.get(Anexo,current['id']);item.file_data=None;db.session.commit()
    request(admin,'post',f'/admin/recuperar-documento/{current["id"]}',400,data={'original':(io.BytesIO(b'adulterado'),'qa.txt')},content_type='multipart/form-data')
    request(admin,'post',f'/admin/recuperar-documento/{current["id"]}',302,data={'original':(io.BytesIO(replacement),'qa.txt')},content_type='multipart/form-data')
    check(request(admin,'get',f'/anexo/{current["id"]}/download').data==replacement,'Recuperação exige hash e restaura original')

    request(admin,'post',f'{path}/acompanhar',302,data={'acompanhar':'1'})
    request(admin,'post',f'{path}/acompanhar',302,data={'acompanhar':'1'})
    request(admin,'get','/home?escopo=acompanhando')
    request(admin,'post','/painel/alertas',302,data={'ativo':'1','antecedencia_dias':'7','escopo':'equipe'})
    request(admin,'post','/painel/alertas',400,data={'antecedencia_dias':'100'})
    request(admin,'post','/relatorios/pesquisas',302,data={'nome_pesquisa':'Pesquisa QA','status':'EM ANÁLISE'})
    saved=row(PesquisaSalva,tenant_id=tenant,nome='Pesquisa QA')
    request(protocolist,'post',f'/relatorios/pesquisas/{saved["id"]}/remover',404)
    request(admin,'post',f'/relatorios/pesquisas/{saved["id"]}/remover',302)
    excel=request(admin,'get','/protocolos/backup/excel')
    check(load_workbook(io.BytesIO(excel.data)).active.max_row>=7,'Excel exporta os protocolos do cliente')
    request(admin,'get','/relatorios/gestao/excel')

    request(admin,'post',f'{path}/tramitar',302,data={'setor_destino_id':str(sector_ids[0]),'destinatario_usuario_id':str(users['alice']['id'])})
    move=row(Movimentacao,protocolo_id=pid,recebido_em=None)
    # Quem ainda não participa não pode nem descobrir este protocolo.
    request(bob,'post',f'{path}/receber',404,data={'movimentacao_id':str(move['id'])})
    request(alice,'post',f'{path}/receber',302,data={'movimentacao_id':str(move['id'])})
    request(alice,'post','/protocolos/atualizar',json={'protocoloId':str(pid),'novoStatus':'EM ANÁLISE'})
    request(alice,'post',f'{path}/tramitar',302,data={'providencia_id':str(move['providencia_id']),
        'setor_destino_id':str(sector_ids[1]),'destinatario_usuario_id':str(users['bob']['id'])})
    request(alice,'post','/protocolos/atualizar',403,json={'protocoloId':str(pid),'novoStatus':'PENDENTE DE DOCUMENTO'})
    move=row(Movimentacao,protocolo_id=pid,recebido_em=None)
    request(bob,'post',f'{path}/receber',302,data={'movimentacao_id':str(move['id'])})
    request(admin,'post',f'{path}/tramitar',302,data={'envio_coletivo':'1','destinos_formato':'campos',
        'destinos_setores':[str(sector_ids[0]),str(sector_ids[2])],
        'destinos_usuarios':[str(users['alice']['id']),str(users['carol']['id'])],
        'destinos_observacoes':['QA primeiro adicional','QA segundo adicional']})
    with app.app_context():
        pending=Movimentacao.query.filter_by(protocolo_id=pid,recebido_em=None).order_by(Movimentacao.id).all()
        received={m.destinatario_usuario_id:(m.id,m.providencia_id) for m in pending}
        check(len([p for p in db.session.get(Protocolo,pid).providencias if not p.concluida_em])==3,'Adicionais mantêm três participações ativas')
    for who,client in [('alice',alice),('carol',carol)]:
        mid,part=received[users[who]['id']]
        request(client,'post',f'{path}/receber',302,data={'movimentacao_id':str(mid)})
        request(client,'post',f'{path}/providencia/{part}/concluir',302,data={'resultado':'QA concluído'})
    request(admin,'post',f'{path}/finalizar',302,data={'resultado':'Conclusão fictícia QA.'})
    request(admin,'post',f'{path}/tramitar',409,data={'setor_destino_id':str(sector_ids[0])})
    request(admin,'post',f'{path}/arquivar',302)
    request(admin,'post',f'{path}/desarquivar',302,data={'justificativa':'Retomar validação fictícia.'})
    check(row(Protocolo,id=pid)['status']=='EM ANÁLISE','Desarquivamento retorna para análise')

    package=request(admin,'get',f'{path}/processo/baixar?formato=zip')
    archive=zipfile.ZipFile(io.BytesIO(package.data));names=archive.namelist()
    check(any(n.endswith('qa-v2.txt') for n in names) and not any(n.endswith('/qa.txt') for n in names),'Pacote inclui apenas anexo atual')
    check(any('despachos/' in n for n in names),'Pacote inclui despacho original')
    check(len(PdfReader(io.BytesIO(archive.read(next(n for n in names if n.startswith('processo-'))))).pages)>=5,'Processo completo tem capa, pedido, histórico e documentos')
    request(viewer,'post',f'{path}/despachos/{dispatch["id"]}/excluir',403,data={'motivo':'Teste inválido','confirmar_exclusao':'1'})
    request(admin,'post',f'{path}/despachos/{dispatch["id"]}/excluir',302,data={'motivo':'Despacho fictício de validação','confirmar_exclusao':'1'})
    request(admin,'get',f'{path}/despachos/{dispatch["id"]}/pdf',404)

    portal=app.test_client()
    request(portal,'post','/portal/qa-a/entrar',302,data={'login':'QA-PORTAL','senha':password})
    request(portal,'get','/portal/meus-dados')
    request(portal,'post','/portal/novo',302,data={'tipo_requerimento':'QA Pedido','requer_ao':'Secretaria QA',
        'observacoes':'Pedido fictício do portal.','declaracao':'on'})
    portal_protocol=row(Protocolo,tenant_id=tenant,modalidade_abertura='remota_requerente')
    check(portal_protocol['status']=='RECEBIDO','Portal sem assinatura tem recebimento automático')
    request(portal,'get',f'/portal/protocolo/{portal_protocol["id"]}')
    request(viewer,'get',f'/portal/protocolo/{portal_protocol["id"]}',(302,403))

    request(admin,'post','/suporte/novo',302,data={'assunto':'Chamado fictício QA','descricao':'Validação de chamado em banco descartável.',
        'categoria':'PROTOCOLOS','prioridade':'NORMAL'})
    ticket=row(ChamadoSuporte,tenant_id=tenant,assunto='Chamado fictício QA')
    request(global_client,'post',f'/suporte/{ticket["id"]}/assumir',302)
    request(global_client,'post',f'/suporte/{ticket["id"]}/status',302,data={'status':'RESOLVIDO'})
    check(row(ChamadoSuporte,id=ticket['id'])['status']=='RESOLVIDO','Chamado assumido e resolvido por administrador geral')
    with app.app_context():check(verify_audit_chain(tenant)[0],'Cadeia de auditoria permanece íntegra')
  output=Path('qa-postgres-tools.json');output.write_text(json.dumps({'passed':len(passed),'cases':passed},ensure_ascii=False,indent=2),encoding='utf-8')
  print(f'PostgreSQL ferramentas: {len(passed)} verificações aprovadas.')
finally:
  with app.app_context():
    db.session.remove();db.drop_all();db.engine.dispose()

